# ERP 재적재 트랙 — 홈택스 기준(세금계산서 + 카드 + 현금영수증) 월별 매출·매입 집계 (2026-10-27).
# 환경변수 SB_URL·SB_KEY 필요(sb.env). 사용: set -a && source sb.env && set +a && python3 hometax_basis_2026.py <출력폴더>
# 계정과목·분개를 전혀 쓰지 않는다. 매입 계산서·카드 지출의 계정 분류 현황 시트만 사용자 확정값(confirmed_account_id)을 그대로 보여준다.
import os, sys, urllib.request, urllib.parse, json, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
U = os.environ['SB_URL']; K = os.environ['SB_KEY']
OUT = sys.argv[1].rstrip('/') + '/'
FROM, TO = '2026-01-01', '2026-09-30'
MONTHS = [f'2026-{m:02d}' for m in range(1, 10)]
ERP_FILE_SALES = [377912660, 686604750, 237037200, 199354220, 203727910, 219279420, 189521600, 191740000, 945819300]  # 사용자 2026 ERP 파일, 부가세 포함
ERP_FILE_COGS = [305111041, 549841718, 198957523, 166550663, 172596581, 185525370, 155755396, 161082287, 766477159]

def fetch_all(path):
    out = []; start = 0
    while True:
        url = U + '/rest/v1/' + urllib.parse.quote(path + '&order=id', safe="=&?.,()*:/'%!")
        r = urllib.request.Request(url, headers={'apikey': K, 'Authorization': 'Bearer ' + K, 'Range': f'{start}-{start + 999}'})
        with urllib.request.urlopen(r) as res: rows = json.loads(res.read())
        out += rows
        if len(rows) < 1000: break
        start += 1000
    return out
M = lambda d: (d or '')[:7]; num = lambda v: float(v or 0)
def mrow(): return {m: 0.0 for m in MONTHS}
accounts = {a['id']: a for a in fetch_all('accounts?select=id,code,name,type')}
aname = lambda i: f"{accounts[i]['code']} {accounts[i]['name']}" if i in accounts else '(미분류)'
ti = fetch_all(f'tax_invoices?select=issue_date,direction,tax_type,supply_amount,tax_amount,total_amount,confirmed_account_id&issue_date=gte.{FROM}&issue_date=lte.{TO}')
cs = fetch_all(f'card_sales?select=tx_date,amount,supply_amount,tax_amount,transaction_type,cancelled_at&tx_date=gte.{FROM}&tx_date=lte.{TO}')
ce = fetch_all(f'card_expenses?select=tx_date,approved_amount,cancel_amount,tax_amount,confirmed_account_id&tx_date=gte.{FROM}&tx_date=lte.{TO}')
cr = fetch_all(f'cash_receipts?select=tx_date,direction,amount,supply_amount,tax_amount,purpose_type&tx_date=gte.{FROM}&tx_date=lte.{TO}')
tx = fetch_all(f'transactions?select=tx_date,amount_in,amount_out,transfer_pair_id&tx_date=gte.{FROM}&tx_date=lte.{TO}')

S = collections.defaultdict(mrow); P = collections.defaultdict(mrow); picls = collections.defaultdict(mrow); cecls = collections.defaultdict(mrow)
for t in ti:
    m = M(t['issue_date'])
    if m not in MONTHS: continue
    kind = '과세' if t['tax_type'] == 'taxable' else '면세'
    D = S if t['direction'] == 'sales' else P
    D[f'계산서 {kind} 공급가액'][m] += num(t['supply_amount']); D[f'계산서 {kind} 세액'][m] += num(t['tax_amount']); D[f'계산서 {kind} 합계'][m] += num(t['total_amount'])
    D[f'계산서 {kind} 건수'][m] += 1
    if t['direction'] != 'sales':
        picls[aname(t['confirmed_account_id'])][m] += num(t['supply_amount'])
for c in cs:
    m = M(c['tx_date'])
    if m not in MONTHS: continue
    sign = -1 if (c['transaction_type'] == 'cancel') else 1
    if c['cancelled_at'] and sign == 1: S['카드 매출 취소표시(참고, 제외)'][m] += num(c['amount']); continue
    S['카드 매출 합계(승인-취소, 부가세 포함)'][m] += sign * num(c['amount'])
    S['카드 매출 건수(승인)'][m] += 1 if sign == 1 else 0
    if num(c['tax_amount']) > 0:
        S['카드 매출 과세 공급가액'][m] += sign * num(c['supply_amount']); S['카드 매출 세액'][m] += sign * num(c['tax_amount'])
    else:
        S['카드 매출 세액 0 금액(면세 또는 미분리)'][m] += sign * num(c['amount'])
for c in ce:
    m = M(c['tx_date'])
    if m not in MONTHS: continue
    amt = num(c['approved_amount']) - num(c['cancel_amount'])
    P['법인카드 매입 합계(승인-취소)'][m] += amt; P['법인카드 매입 세액(카드사 표기)'][m] += num(c['tax_amount']); P['법인카드 매입 건수'][m] += 1
    cecls[aname(c['confirmed_account_id'])][m] += amt
for r in cr:
    m = M(r['tx_date'])
    if m not in MONTHS: continue
    D = S if r['direction'] == 'sales' else P
    D['현금영수증 공급가액'][m] += num(r['supply_amount']); D['현금영수증 세액'][m] += num(r['tax_amount']); D['현금영수증 합계'][m] += num(r['amount']); D['현금영수증 건수'][m] += 1
dep = mrow(); wd = mrow()
for t in tx:
    m = M(t['tx_date'])
    if m in dep and not t['transfer_pair_id']: dep[m] += num(t['amount_in']); wd[m] += num(t['amount_out'])
for m in MONTHS:
    S['매출 합계(부가세 포함) = 계산서+카드+현금영수증'][m] = S['계산서 과세 합계'][m] + S['계산서 면세 합계'][m] + S['카드 매출 합계(승인-취소, 부가세 포함)'][m] + S['현금영수증 합계'][m]
    S['매출 공급가액 합계(세액 제외)'][m] = S['계산서 과세 공급가액'][m] + S['계산서 면세 공급가액'][m] + S['카드 매출 과세 공급가액'][m] + S['카드 매출 세액 0 금액(면세 또는 미분리)'][m] + S['현금영수증 공급가액'][m]
    S['매출 세액 합계'][m] = S['계산서 과세 세액'][m] + S['카드 매출 세액'][m] + S['현금영수증 세액'][m]
    P['매입 합계(부가세 포함) = 계산서+법인카드+현금영수증'][m] = P['계산서 과세 합계'][m] + P['계산서 면세 합계'][m] + P['법인카드 매입 합계(승인-취소)'][m] + P['현금영수증 합계'][m]
    S['참고: ERP 매출(파일, 부가세 포함)'][m] = ERP_FILE_SALES[MONTHS.index(m)]
    S['참고: 커버리지 = 매출 합계 / ERP 매출'][m] = S['매출 합계(부가세 포함) = 계산서+카드+현금영수증'][m] / ERP_FILE_SALES[MONTHS.index(m)]
    S['참고: 통장 입금(이체 제외)'][m] = dep[m]
    P['참고: ERP 매입원가(파일)'][m] = ERP_FILE_COGS[MONTHS.index(m)]
    P['참고: 통장 출금(이체 제외)'][m] = wd[m]

wb = Workbook(); HF = PatternFill('solid', fgColor='DDE3EA'); B = Font(bold=True)
def section(ws, r, title, d, order, bold=(), pct=()):
    ws.cell(r, 1, title).font = B; ws.cell(r, 1).fill = HF
    for j, m in enumerate(MONTHS, 2): c = ws.cell(r, j, f'{int(m[5:])}월'); c.font = B; c.fill = HF
    ws.cell(r, 11, '합계').font = B; ws.cell(r, 11).fill = HF; r += 1
    for k in order:
        if k not in d: continue
        ws.cell(r, 1, k).font = Font(bold=k in bold)
        for j, m in enumerate(MONTHS, 2):
            c = ws.cell(r, j, round(d[k][m], 3) if k in pct else round(d[k][m])); c.number_format = '0%' if k in pct else '#,##0'
        if k not in pct:
            c = ws.cell(r, 11, round(sum(d[k].values()))); c.number_format = '#,##0'; c.font = B
        elif k.startswith('참고: 커버리지'):
            tot = sum(S['매출 합계(부가세 포함) = 계산서+카드+현금영수증'].values()) / sum(ERP_FILE_SALES)
            c = ws.cell(r, 11, round(tot, 3)); c.number_format = '0%'; c.font = B
        r += 1
    ws.column_dimensions['A'].width = 50
    for j in range(2, 12): ws.column_dimensions[get_column_letter(j)].width = 14
    return r + 1
ws = wb.active; ws.title = '매출(홈택스 기준)'
ws.cell(1, 1, '홈택스 기준 매출 = 매출 세금계산서(발행일) + 카드 매출(승인일, 승인-취소) + 현금영수증(발행일). 계정과목·분개 사용 없음. 추출 2026-10-27').font = Font(bold=True)
ws.cell(2, 1, '현금영수증 39건 = 전량(2026-10-27 사용자 확인). 카드 매출은 카드사 자료의 공급가·세액을 그대로 사용(1,208건 분리, 319건은 세액 0 = 면세 또는 미분리 — 별도 행).').font = Font(italic=True, color='555555')
section(ws, 4, '매출', S, ['매출 합계(부가세 포함) = 계산서+카드+현금영수증', '매출 공급가액 합계(세액 제외)', '매출 세액 합계',
    '계산서 과세 공급가액', '계산서 과세 세액', '계산서 과세 합계', '계산서 과세 건수', '계산서 면세 공급가액', '계산서 면세 합계', '계산서 면세 건수',
    '카드 매출 합계(승인-취소, 부가세 포함)', '카드 매출 과세 공급가액', '카드 매출 세액', '카드 매출 세액 0 금액(면세 또는 미분리)', '카드 매출 건수(승인)', '카드 매출 취소표시(참고, 제외)',
    '현금영수증 공급가액', '현금영수증 세액', '현금영수증 합계', '현금영수증 건수',
    '참고: ERP 매출(파일, 부가세 포함)', '참고: 커버리지 = 매출 합계 / ERP 매출', '참고: 통장 입금(이체 제외)'],
    bold={'매출 합계(부가세 포함) = 계산서+카드+현금영수증'}, pct={'참고: 커버리지 = 매출 합계 / ERP 매출'})
ws = wb.create_sheet('매입(홈택스 기준)')
ws.cell(1, 1, '홈택스 기준 매입 = 매입 세금계산서(발행일) + 법인카드 사용(승인-취소) + 현금영수증(지출증빙). 상품매입과 비용성 매입이 섞여 있다 — 분리는 다음 시트의 계정 분류 현황 참고.').font = Font(bold=True)
section(ws, 3, '매입', P, ['매입 합계(부가세 포함) = 계산서+법인카드+현금영수증',
    '계산서 과세 공급가액', '계산서 과세 세액', '계산서 과세 합계', '계산서 과세 건수', '계산서 면세 공급가액', '계산서 면세 합계', '계산서 면세 건수',
    '법인카드 매입 합계(승인-취소)', '법인카드 매입 세액(카드사 표기)', '법인카드 매입 건수', '현금영수증 공급가액', '현금영수증 세액', '현금영수증 합계', '현금영수증 건수',
    '참고: ERP 매입원가(파일)', '참고: 통장 출금(이체 제외)'], bold={'매입 합계(부가세 포함) = 계산서+법인카드+현금영수증'})
ws = wb.create_sheet('매입계산서 계정별(확정)')
ws.cell(1, 1, '매입 계산서 공급가액을 사용자가 확정한 계정과목별로. (미분류)가 크면 그 달은 상품매입/비용 분리가 안 된다. 임의 배정 없음.').font = Font(italic=True, color='555555')
section(ws, 3, '매입 계산서 공급가액', picls, sorted(picls, key=lambda k: (k == '(미분류)', k)), bold={'(미분류)'})
ws = wb.create_sheet('법인카드 계정별(확정)')
ws.cell(1, 1, '법인카드 사용액(승인-취소)을 사용자가 확정한 계정과목별로. 임의 배정 없음.').font = Font(italic=True, color='555555')
section(ws, 3, '법인카드 사용액', cecls, sorted(cecls, key=lambda k: (k == '(미분류)', k)), bold={'(미분류)'})
wb.save(OUT + '2026_1-9월_홈택스기준_매출매입_2026-10-27.xlsx')
print('saved')
for k in ['매출 합계(부가세 포함) = 계산서+카드+현금영수증', '참고: 커버리지 = 매출 합계 / ERP 매출']: print(k, [round(S[k][m] / 1e6, 1) if '커버' not in k else round(S[k][m], 2) for m in MONTHS])
print('매입 합계', [round(P['매입 합계(부가세 포함) = 계산서+법인카드+현금영수증'][m] / 1e6, 1) for m in MONTHS])
print('매입계산서 미분류', [round(picls['(미분류)'][m] / 1e6, 1) for m in MONTHS])
print('법인카드 미분류', [round(cecls['(미분류)'][m] / 1e6, 1) for m in MONTHS])
