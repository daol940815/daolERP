# ERP 재적재 트랙 — 2026년 월별 집계(원본 합계·분류 완료율·계정별·분개 손익) 엑셀. 환경변수 SB_URL·SB_KEY 필요(sb.env). docs/erp-reload-track.md 참고.
# 사용: set -a && source sb.env && set +a && python3 monthly_2026.py   (OUT 경로는 아래 상수 수정)
import os, urllib.request, urllib.parse, json, collections, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
U = os.environ['SB_URL']; K = os.environ['SB_KEY']
OUT = '/tmp/claude-0/-home-user-daolERP/d1f2072d-70cb-5d40-986e-e5b4cb41c248/scratchpad/review/'
FROM, TO = '2026-01-01', '2026-09-30'
MONTHS = [f'2026-{m:02d}' for m in range(1, 10)]

def fetch_all(path):
    out = []; start = 0
    while True:
        url = U + '/rest/v1/' + urllib.parse.quote(path + ('' if 'order=' in path else '&order=id'), safe="=&?.,()*:/'%!")
        r = urllib.request.Request(url, headers={'apikey': K, 'Authorization': 'Bearer ' + K, 'Range': f'{start}-{start + 999}'})
        with urllib.request.urlopen(r) as res:
            rows = json.loads(res.read())
        out += rows
        if len(rows) < 1000: break
        start += 1000
    return out

accounts = {a['id']: a for a in fetch_all('accounts?select=id,code,name,type,side_on_in,side_on_out&order=code')}
banks = {b['id']: b for b in fetch_all('bank_accounts?select=id,alias,bank_name,account_type')}
print('accounts', len(accounts), collections.Counter(a['type'] for a in accounts.values()))
orders = fetch_all(f'erp_orders?select=id,order_date,total_amount,collect_status,source,outstanding_amount&order_date=gte.{FROM}&order_date=lte.{TO}')
oid = {o['id']: o for o in orders}
items = [i for i in fetch_all('erp_order_items?select=order_id,line_total,purchase_total,is_canceled,is_vip,is_prepayment,order_kind,item_name,shipping_fee') if i['order_id'] in oid]
tis = fetch_all(f'tax_invoices?select=direction,issue_date,supply_amount,tax_amount,total_amount,tax_type,confirmed_account_id,payment_status,matched_transaction_id&issue_date=gte.{FROM}&issue_date=lte.{TO}')
txs = fetch_all(f'transactions?select=tx_date,amount_in,amount_out,confirmed_account_id,suggested_account_id,transfer_pair_id,status,tx_kind,bank_account_id,vendor_id&tx_date=gte.{FROM}&tx_date=lte.{TO}')
cs = fetch_all(f'card_sales?select=tx_date,amount,supply_amount,tax_amount,transaction_type,cancelled_at&tx_date=gte.{FROM}&tx_date=lte.{TO}')
ce = fetch_all(f'card_expenses?select=tx_date,approved_amount,cancel_amount,confirmed_account_id,suggested_account_id,classify_status,usage_type&tx_date=gte.{FROM}&tx_date=lte.{TO}')
cr = fetch_all(f'cash_receipts?select=tx_date,amount,direction,transaction_type&tx_date=gte.{FROM}&tx_date=lte.{TO}')
jl = fetch_all(f'journal_lines?select=account_id,amount,side,journal_entries!inner(entry_date,entry_type,source_type)&journal_entries.entry_date=gte.{FROM}&journal_entries.entry_date=lte.{TO}')
print('orders', len(orders), 'items', len(items), 'tax', len(tis), 'tx', len(txs), 'cs', len(cs), 'ce', len(ce), 'cr', len(cr), 'jl', len(jl))
print('tx status', collections.Counter(t['status'] for t in txs), 'tx_kind', collections.Counter(t['tx_kind'] for t in txs))
print('ce classify_status', collections.Counter(c['classify_status'] for c in ce), 'usage_type', collections.Counter(c['usage_type'] for c in ce))
print('cs transaction_type', collections.Counter(c['transaction_type'] for c in cs))
print('ti payment_status', collections.Counter(t['payment_status'] for t in tis), 'direction', collections.Counter(t['direction'] for t in tis))
print('je entry_type', collections.Counter(l['journal_entries']['entry_type'] for l in jl), 'source', collections.Counter(l['journal_entries']['source_type'] for l in jl))

M = lambda d: (d or '')[:7]
def mrow(): return {m: 0 for m in MONTHS}
num = lambda v: float(v or 0)

# ── ERP ──
erp = collections.defaultdict(mrow)
for i in items:
    m = M(oid[i['order_id']]['order_date'])
    if m not in erp['_']: continue
    if i['is_canceled']: erp['취소 금액(참고)'][m] += num(i['line_total']); continue
    if i['is_vip']: erp['VIP(매출 제외, 참고)'][m] += num(i['line_total']); continue
    if i['is_prepayment']: erp['선결제(매출 제외, 참고)'][m] += num(i['line_total']); continue
    kind = i['order_kind'] or '(없음)'
    erp[f'매출 - {kind}'][m] += num(i['line_total'])
    erp['매출 합계'][m] += num(i['line_total'])
    erp['매입원가 합계(품목 매입금액)'][m] += num(i['purchase_total'])
    erp['배송비(매출 포함분)'][m] += num(i['shipping_fee'])
erp.pop('_', None)
for m in MONTHS: erp['매출총이익(매출-매입원가)'][m] = erp['매출 합계'][m] - erp['매입원가 합계(품목 매입금액)'][m]
ocnt = collections.defaultdict(mrow)
for o in orders:
    m = M(o['order_date'])
    if m in ocnt['주문 건수']: ocnt['주문 건수'][m] += 1; ocnt['미수 잔액(주문 기준)'][m] += num(o['outstanding_amount'])

# ── 계산서 ──
ti = collections.defaultdict(mrow)
for t in tis:
    m = M(t['issue_date']); d = '매출' if t['direction'] in ('sales', 'out', '매출') else '매입'
    if m not in ti['_']: continue
    ti[f'{d} 계산서 공급가액'][m] += num(t['supply_amount']); ti[f'{d} 계산서 세액'][m] += num(t['tax_amount']); ti[f'{d} 계산서 합계'][m] += num(t['total_amount'])
    ti[f'{d} 계산서 건수'][m] += 1
    if t['matched_transaction_id'] or t['payment_status'] in ('paid', 'matched', 'completed'): ti[f'{d} 계산서 입출금 매칭 건수'][m] += 1
ti.pop('_', None)
print('direction values', collections.Counter(t['direction'] for t in tis))

# ── 통장 ──
tx = collections.defaultdict(mrow); txacc_in = collections.defaultdict(mrow); txacc_out = collections.defaultdict(mrow)
txacc_in_s = collections.defaultdict(mrow); txacc_out_s = collections.defaultdict(mrow)
cov = collections.defaultdict(mrow)
aname = lambda i: f"{accounts[i]['code']} {accounts[i]['name']}" if i in accounts else '(미분류)'
for t in txs:
    m = M(t['tx_date'])
    if m not in tx['_']: continue
    ain, aout = num(t['amount_in']), num(t['amount_out'])
    cov['통장 건수'][m] += 1
    if t['transfer_pair_id']:
        tx['내부 이체 입금(제외)'][m] += ain; tx['내부 이체 출금(제외)'][m] += aout; cov['이체쌍 건수'][m] += 1; continue
    tx['입금 합계(이체 제외)'][m] += ain; tx['출금 합계(이체 제외)'][m] += aout
    cov['분류 대상 건수(이체 제외)'][m] += 1
    cov['분류 대상 금액(입금+출금)'][m] += ain + aout
    acc = t['confirmed_account_id']
    if acc: cov['계정 확정 건수'][m] += 1; cov['계정 확정 금액'][m] += ain + aout
    if t['status'] == 'confirmed': cov['상태 confirmed 건수'][m] += 1
    eff = acc or t['suggested_account_id']
    if eff: cov['확정+추천 건수'][m] += 1
    (txacc_in if ain else txacc_out)[aname(acc)][m] += ain or aout
    (txacc_in_s if ain else txacc_out_s)[aname(eff)][m] += ain or aout
tx.pop('_', None)
for m in MONTHS:
    n = cov['분류 대상 건수(이체 제외)'][m]; a = cov['분류 대상 금액(입금+출금)'][m]
    cov['계정 확정 비율(건수)'][m] = round(cov['계정 확정 건수'][m] / n, 3) if n else 0
    cov['계정 확정 비율(금액)'][m] = round(cov['계정 확정 금액'][m] / a, 3) if a else 0
    cov['확정+추천 비율(건수)'][m] = round(cov['확정+추천 건수'][m] / n, 3) if n else 0

# ── 카드 ──
cd = collections.defaultdict(mrow); ceacc = collections.defaultdict(mrow); ccov = collections.defaultdict(mrow)
for c in cs:
    m = M(c['tx_date'])
    if m not in cd['_']: continue
    amt = num(c['amount'])
    if c['cancelled_at'] or (c['transaction_type'] or '').lower() in ('cancel', '취소', 'refund'): cd['카드 매출 취소(참고)'][m] += amt
    else: cd['카드 매출'][m] += amt
for c in ce:
    m = M(c['tx_date'])
    if m not in cd['_']: continue
    amt = num(c['approved_amount']) - num(c['cancel_amount'])
    cd['카드 지출(승인-취소)'][m] += amt; ccov['카드 지출 건수'][m] += 1
    if c['confirmed_account_id']: ccov['계정 확정 건수'][m] += 1
    ceacc[aname(c['confirmed_account_id'])][m] += amt
for r in cr:
    m = M(r['tx_date'])
    if m in cd['_']: cd[f"현금영수증 {r['direction']}"][m] += num(r['amount'])
cd.pop('_', None)
for m in MONTHS:
    n = ccov['카드 지출 건수'][m]; ccov['계정 확정 비율'][m] = round(ccov['계정 확정 건수'][m] / n, 3) if n else 0

# ── 분개 손익 ──
pl = collections.defaultdict(mrow); plcnt = mrow()
for l in jl:
    m = M(l['journal_entries']['entry_date'])
    if m not in plcnt: continue
    a = accounts.get(l['account_id']); t = (a or {}).get('type', '?')
    amt = num(l['amount']); side = l['side']
    if t in ('revenue', 'income', '수익'):
        pl[f"[수익] {aname(l['account_id'])}"][m] += amt if side == 'credit' else -amt
    elif t in ('expense', 'cost', '비용'):
        pl[f"[비용] {aname(l['account_id'])}"][m] += amt if side == 'debit' else -amt
    plcnt[m] += 1
for m in MONTHS:
    pl['수익 합계'][m] = sum(v[m] for k, v in pl.items() if k.startswith('[수익]'))
    pl['비용 합계'][m] = sum(v[m] for k, v in pl.items() if k.startswith('[비용]'))
    pl['분개 기준 손익'][m] = pl['수익 합계'][m] - pl['비용 합계'][m]
    pl['분개 라인 수'][m] = plcnt[m]

# ── 엑셀 ──
wb = Workbook(); HF = PatternFill('solid', fgColor='DDE3EA'); B = Font(bold=True)
def put(ws, r, label, row, bold=False, pct=False):
    ws.cell(r, 1, label).font = Font(bold=bold)
    for j, m in enumerate(MONTHS, 2):
        c = ws.cell(r, j, row[m]); c.number_format = '0.0%' if pct else '#,##0'
        if bold: c.font = B
    if not pct:
        c = ws.cell(r, 11, sum(row[m] for m in MONTHS)); c.number_format = '#,##0'; c.font = B
def header(ws, r, title):
    ws.cell(r, 1, title).font = B; ws.cell(r, 1).fill = HF
    for j, m in enumerate(MONTHS, 2): c = ws.cell(r, j, f'{int(m[5:])}월'); c.font = B; c.fill = HF
    ws.cell(r, 11, '합계').font = B; ws.cell(r, 11).fill = HF
    ws.column_dimensions['A'].width = 44
    for j in range(2, 12): ws.column_dimensions[get_column_letter(j)].width = 14
def section(ws, r, title, d, order=None, bold_keys=(), pct_keys=()):
    header(ws, r, title); r += 1
    keys = order or sorted(d.keys())
    for k in keys:
        if k in d: put(ws, r, k, d[k], bold=k in bold_keys, pct=k in pct_keys); r += 1
    return r + 1

ws = wb.active; ws.title = '요약'
ws.cell(1, 1, '2026년 1~9월 월별 집계 — 원본 기준 (계정과목 불필요)   추출 2026-10-26, 계정과목 미기재와 무관한 합계').font = Font(bold=True, size=12)
r = 3
r = section(ws, r, 'ERP 주문 (주문일 기준, 취소·VIP·선결제 제외)', erp,
            order=['매출 합계'] + sorted(k for k in erp if k.startswith('매출 - ')) + ['배송비(매출 포함분)', '매입원가 합계(품목 매입금액)', '매출총이익(매출-매입원가)', '취소 금액(참고)', 'VIP(매출 제외, 참고)', '선결제(매출 제외, 참고)'],
            bold_keys={'매출 합계', '매출총이익(매출-매입원가)'})
r = section(ws, r, 'ERP 주문 건수·미수', ocnt, order=['주문 건수', '미수 잔액(주문 기준)'])
r = section(ws, r, '세금계산서 (발행일 기준)', ti,
            order=['매출 계산서 공급가액', '매출 계산서 세액', '매출 계산서 합계', '매출 계산서 건수', '매출 계산서 입출금 매칭 건수', '매입 계산서 공급가액', '매입 계산서 세액', '매입 계산서 합계', '매입 계산서 건수', '매입 계산서 입출금 매칭 건수'],
            bold_keys={'매출 계산서 공급가액', '매입 계산서 공급가액'})
r = section(ws, r, '통장 (거래일 기준, 내부 이체쌍 제외)', tx, order=['입금 합계(이체 제외)', '출금 합계(이체 제외)', '내부 이체 입금(제외)', '내부 이체 출금(제외)'], bold_keys={'입금 합계(이체 제외)', '출금 합계(이체 제외)'})
r = section(ws, r, '카드·현금영수증', cd, bold_keys={'카드 매출', '카드 지출(승인-취소)'})
r = section(ws, r, '교차 확인', {
    'ERP 매출 - 매출 계산서 공급가액': {m: erp['매출 합계'][m] - ti['매출 계산서 공급가액'][m] for m in MONTHS},
    '통장 입금 - 매출 계산서 합계 - 카드 매출': {m: tx['입금 합계(이체 제외)'][m] - ti['매출 계산서 합계'][m] - cd['카드 매출'][m] for m in MONTHS},
    'ERP 매입원가 - 매입 계산서 공급가액': {m: erp['매입원가 합계(품목 매입금액)'][m] - ti['매입 계산서 공급가액'][m] for m in MONTHS},
}, order=['ERP 매출 - 매출 계산서 공급가액', '통장 입금 - 매출 계산서 합계 - 카드 매출', 'ERP 매입원가 - 매입 계산서 공급가액'])
ws.cell(r, 1, '주의: 교차 확인 차액은 시차·무계산서 매출·합산 발행·선결제·부가세 포함 여부 때문에 0이 되지 않는다. 추세 참고용.').font = Font(italic=True, color='555555')

ws = wb.create_sheet('분류 완료율')
ws.cell(1, 1, '계정과목 기재 현황 — 손익(계정과목별) 숫자를 얼마나 믿을 수 있는지의 척도').font = Font(bold=True, size=12)
r = 3
r = section(ws, r, '통장 (이체쌍 제외)', cov, order=['통장 건수', '이체쌍 건수', '분류 대상 건수(이체 제외)', '계정 확정 건수', '계정 확정 비율(건수)', '분류 대상 금액(입금+출금)', '계정 확정 금액', '계정 확정 비율(금액)', '확정+추천 건수', '확정+추천 비율(건수)', '상태 confirmed 건수'],
            pct_keys={'계정 확정 비율(건수)', '계정 확정 비율(금액)', '확정+추천 비율(건수)'})
r = section(ws, r, '카드 지출', ccov, order=['카드 지출 건수', '계정 확정 건수', '계정 확정 비율'], pct_keys={'계정 확정 비율'})
r = section(ws, r, '분개', {'분개 라인 수': pl['분개 라인 수']}, order=['분개 라인 수'])

def acc_sheet(title, din, dout, note):
    ws = wb.create_sheet(title); ws.cell(1, 1, note).font = Font(italic=True, color='555555'); r = 3
    r = section(ws, r, '입금 계정과목별', din, order=sorted(din, key=lambda k: (k == '(미분류)', k)), bold_keys={'(미분류)'})
    r = section(ws, r, '출금 계정과목별', dout, order=sorted(dout, key=lambda k: (k == '(미분류)', k)), bold_keys={'(미분류)'})
acc_sheet('통장 계정별(확정)', txacc_in, txacc_out, '사용자가 확정한 계정과목만 반영. (미분류) 행이 크면 그 달의 계정별 숫자는 불완전.')
acc_sheet('통장 계정별(확정+추천)', txacc_in_s, txacc_out_s, '확정이 없으면 시스템 추천 계정을 적용 — 참고용(추천은 확정이 아님).')
ws = wb.create_sheet('카드지출 계정별'); r = 3
ws.cell(1, 1, '카드 지출 확정 계정과목별 (승인-취소).').font = Font(italic=True, color='555555')
section(ws, r, '카드 지출 계정과목별', ceacc, order=sorted(ceacc, key=lambda k: (k == '(미분류)', k)), bold_keys={'(미분류)'})
ws = wb.create_sheet('분개 기준 손익'); r = 3
ws.cell(1, 1, '분개(journal) 기준 수익·비용 — 분개가 생성된 범위만 반영되므로 분류 완료율과 함께 볼 것.').font = Font(italic=True, color='555555')
section(ws, r, '분개 기준', pl, order=['수익 합계'] + sorted(k for k in pl if k.startswith('[수익]')) + ['비용 합계'] + sorted(k for k in pl if k.startswith('[비용]')) + ['분개 기준 손익', '분개 라인 수'],
        bold_keys={'수익 합계', '비용 합계', '분개 기준 손익'})
wb.save(OUT + '2026_1-9월_월별집계_2026-10-26.xlsx')
print('saved')
for k in ['매출 합계', '매입원가 합계(품목 매입금액)']: print(k, [int(erp[k][m]) for m in MONTHS])
print('계산서 매출', [int(ti['매출 계산서 공급가액'][m]) for m in MONTHS])
print('입금', [int(tx['입금 합계(이체 제외)'][m]) for m in MONTHS])
print('출금', [int(tx['출금 합계(이체 제외)'][m]) for m in MONTHS])
print('확정비율', [cov['계정 확정 비율(건수)'][m] for m in MONTHS])
print('카드지출 확정비율', [ccov['계정 확정 비율'][m] for m in MONTHS])
print('분개 라인', [plcnt[m] for m in MONTHS])
