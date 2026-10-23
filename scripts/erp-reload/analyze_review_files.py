# ERP 재적재 트랙 — 검수 파일 1차 분석 (2026-10-21 총괄). docs/erp-reload-track.md 참고.
# 사용: python3 -I analyze_review_files.py <업로드폴더> <출력폴더>
#   업로드폴더: 2025 ERP·2026 ERP(판매내역 양식 27열)·고객관리문서 xlsx 3개 (파일명은 아래 FILES 수정)
#   출력: summary.json (오탈자 후보·문서 미일치·채널 불일치·구분 분포)
import sys
import openpyxl, re, json, sys, collections, difflib, os, urllib.request, urllib.parse
FILES_ERP = ['0171b981-2025_ERP_____.xlsx', '5e6e20ac-2026_ERP_____.xlsx']  # 사용자 업로드 파일명으로 수정
FILE_DOC = 'cb9d73b3-2026___________25_.xlsx'
UP = sys.argv[1].rstrip('/') + '/'
OUT = sys.argv[2].rstrip('/') + '/'

def s(v):
    if v is None: return ''
    return str(v).strip()

# ── ERP rows ──
erp = []
for fn in [*FILES_ERP]:
    wb = openpyxl.load_workbook(UP + fn, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        hdr = [s(c) for c in next(rows)]
        for r in rows:
            if all(c is None for c in r): continue
            d = dict(zip(hdr[:27], list(r)[:27]))
            # duplicate header names: 합계금액 x2, 비고 x3 -> keep positional too
            d['_sheet'] = ws.title
            d['_bank'] = s(r[4]); d['_branch'] = s(r[5]); d['_mgr'] = s(r[6])
            d['_kind'] = s(r[11]); d['_channel'] = s(r[21]); d['_writer'] = s(r[22]); d['_gubun'] = s(r[23])
            d['_status'] = s(r[0]); d['_reg'] = s(r[24]); d['_item'] = s(r[10]); d['_amt'] = r[16]
            d['_ymd'] = f"{s(r[1])}-{s(r[2]).zfill(2)}-{s(r[3]).zfill(2)}"
            erp.append(d)
print('ERP rows', len(erp))
C = collections.Counter
print('sheets', C(d['_sheet'] for d in erp).most_common()[:3], '...')
print('date range', min(d['_ymd'] for d in erp), max(d['_ymd'] for d in erp))
print('상태', C(d['_status'] for d in erp).most_common())
print('구분', C(d['_gubun'] for d in erp).most_common())
print('샘플,지점,개별', C(d['_kind'] for d in erp).most_common())
print('채널', C(d['_channel'] for d in erp).most_common())
print('작성자', C(d['_writer'] for d in erp).most_common())
print('은행명 distinct', len(set(d['_bank'] for d in erp)))
print('은행명 top', C(d['_bank'] for d in erp).most_common(40))
branches = set((d['_bank'], d['_branch']) for d in erp)
mgrs = set((d['_bank'], d['_branch'], d['_mgr']) for d in erp)
print('은행|지점 distinct', len(branches), ' 은행|지점|담당자 distinct', len(mgrs))
print('빈 은행명', sum(1 for d in erp if not d['_bank']), '빈 지점명', sum(1 for d in erp if not d['_branch']), '빈 담당자', sum(1 for d in erp if not d['_mgr']))

# ── 고객관리문서 ──
doc = {}
docrows = []
wb = openpyxl.load_workbook(UP + FILE_DOC, read_only=True, data_only=True)
for ws in wb.worksheets:
    if ws.title == '집계': continue
    started = False
    for r in ws.iter_rows(values_only=True):
        if not started:
            if s(r[1]) == '은행명' and s(r[2]) == '지점명': started = True
            continue
        if not s(r[1]) and not s(r[2]) and not s(r[3]): continue
        row = {'sheet': ws.title, 'bank': s(r[1]), 'branch': s(r[2]), 'mgr': s(r[3]), 'daol': s(r[4]), 'prev': s(r[5]),
               'first': s(r[6])[:10], 'last': s(r[7])[:10], 'move': s(r[8]), 'memo': s(r[9]), 'type': s(r[47])}
        docrows.append(row)
        doc[(row['bank'], row['branch'], row['mgr'])] = row
print('문서 행', len(docrows), 'distinct key', len(doc))
print('문서 다올 담당', C(r['daol'] for r in docrows).most_common())
print('문서 은행명', C(r['bank'] for r in docrows).most_common(30))
dupkeys = C((r['bank'], r['branch'], r['mgr']) for r in docrows)
print('문서 내 중복 키', [(k, n) for k, n in dupkeys.items() if n > 1][:20], '...total', sum(1 for n in dupkeys.values() if n > 1))

# ── 2) ERP combo vs doc ──
combo_cnt = C((d['_bank'], d['_branch'], d['_mgr']) for d in erp)
in_doc = sum(n for k, n in combo_cnt.items() if k in doc)
not_in_doc = [(k, n) for k, n in combo_cnt.items() if k not in doc]
print('ERP 조합 수', len(combo_cnt), '문서 일치 조합', sum(1 for k in combo_cnt if k in doc), '행 기준 일치', in_doc, '/', len(erp))
# partial: bank+branch in doc but mgr not
bb_doc = set((r['bank'], r['branch']) for r in docrows)
bb_only = [(k, n) for k, n in not_in_doc if (k[0], k[1]) in bb_doc]
bank_doc = set(r['bank'] for r in docrows)
bank_only = [(k, n) for k, n in not_in_doc if (k[0], k[1]) not in bb_doc and k[0] in bank_doc]
none = [(k, n) for k, n in not_in_doc if k[0] not in bank_doc]
print(' 문서 미일치: 지점은 있고 담당자 없음', len(bb_only), '행', sum(n for _, n in bb_only))
print(' 문서 미일치: 은행은 있고 지점 없음', len(bank_only), '행', sum(n for _, n in bank_only))
print(' 문서 미일치: 은행도 없음', len(none), '행', sum(n for _, n in none))
print(' 은행도 없음 top', sorted(none, key=lambda x: -x[1])[:15])

# channel vs daol
ch_match = ch_diff = ch_nodoc = 0
diff_pairs = C()
for d in erp:
    k = (d['_bank'], d['_branch'], d['_mgr'])
    if k in doc:
        if doc[k]['daol'] == d['_channel']: ch_match += 1
        else:
            ch_diff += 1; diff_pairs[(d['_channel'], doc[k]['daol'])] += 1
    else: ch_nodoc += 1
print('채널 vs 다올담당: 일치', ch_match, '불일치', ch_diff, '문서없음', ch_nodoc)
print(' 불일치 쌍 top', diff_pairs.most_common(20))

# ── 1) typo candidates ──
def norm(x):
    x = x.lower()
    x = re.sub(r'\s+', '', x)
    x = re.sub(r'(님|지점|점|센터|영업부|부)$', '', x)
    return x
def mnorm(x):
    x = re.sub(r'\s+', '', x)
    x = re.sub(r'님$', '', x)
    return x

# bank variants
bank_cnt = C(d['_bank'] for d in erp)
groups = collections.defaultdict(list)
for b in bank_cnt: groups[norm(b)].append(b)
bank_var = {k: v for k, v in groups.items() if len(v) > 1}
print('은행명 정규화 후 변형 그룹', len(bank_var), list(bank_var.items())[:15])
# fuzzy among banks
banks = list(bank_cnt)
fz = []
for i in range(len(banks)):
    for j in range(i + 1, len(banks)):
        a, b = banks[i], banks[j]
        if norm(a) == norm(b): continue
        r = difflib.SequenceMatcher(None, a, b).ratio()
        if r >= 0.8 and min(len(a), len(b)) >= 3: fz.append((a, bank_cnt[a], b, bank_cnt[b], round(r, 2)))
print('은행명 유사 쌍', len(fz), sorted(fz, key=lambda x: -x[4])[:25])

# branch variants within same bank
br_cnt = C((d['_bank'], d['_branch']) for d in erp)
bybank = collections.defaultdict(list)
for (b, br) in br_cnt: bybank[b].append(br)
br_fz = []
br_norm_groups = []
for b, brs in bybank.items():
    g = collections.defaultdict(list)
    for br in brs: g[norm(br)].append(br)
    for k, v in g.items():
        if len(v) > 1: br_norm_groups.append((b, [(x, br_cnt[(b, x)]) for x in v]))
    for i in range(len(brs)):
        for j in range(i + 1, len(brs)):
            a, c = brs[i], brs[j]
            if not a or not c or norm(a) == norm(c): continue
            r = difflib.SequenceMatcher(None, a, c).ratio()
            if r >= 0.8 and min(len(a), len(c)) >= 3: br_fz.append((b, a, br_cnt[(b, a)], c, br_cnt[(b, c)], round(r, 2)))
print('지점명 정규화 변형 그룹', len(br_norm_groups), br_norm_groups[:15])
print('지점명 유사 쌍', len(br_fz), sorted(br_fz, key=lambda x: -x[5])[:25])

# manager variants within same bank+branch
mg_cnt = C((d['_bank'], d['_branch'], d['_mgr']) for d in erp)
bybb = collections.defaultdict(list)
for (b, br, m) in mg_cnt: bybb[(b, br)].append(m)
mg_norm_groups = []; mg_fz = []
for bb, ms in bybb.items():
    g = collections.defaultdict(list)
    for m in ms: g[mnorm(m)].append(m)
    for k, v in g.items():
        if len(v) > 1: mg_norm_groups.append((bb, [(x, mg_cnt[bb + (x,)]) for x in v]))
    for i in range(len(ms)):
        for j in range(i + 1, len(ms)):
            a, c = ms[i], ms[j]
            if not a or not c or mnorm(a) == mnorm(c): continue
            # same surname+name different title, or 1-char diff
            na, nc = re.sub(r'\s.*$', '', a), re.sub(r'\s.*$', '', c)
            r = difflib.SequenceMatcher(None, a, c).ratio()
            if na == nc or r >= 0.75: mg_fz.append((bb, a, mg_cnt[bb + (a,)], c, mg_cnt[bb + (c,)], round(r, 2)))
print('담당자 정규화 변형 그룹', len(mg_norm_groups), mg_norm_groups[:15])
print('담당자 유사 쌍', len(mg_fz), sorted(mg_fz, key=lambda x: -x[5])[:30])
# honorific/title patterns in manager names
print('담당자 끝 패턴', C(re.sub(r'^.*?(\S+)$', r'\1', m)[-3:] for (_, _, m) in mg_cnt if m).most_common(15))
no_nim = [(k, n) for k, n in mg_cnt.items() if k[2] and not k[2].endswith('님')]
print('님 없는 담당자 조합', len(no_nim), '행', sum(n for _, n in no_nim), sorted(no_nim, key=lambda x: -x[1])[:20])

json.dump({'erp_n': len(erp), 'docrows': docrows, 'not_in_doc': [(list(k), n) for k, n in not_in_doc],
           'bank_fz': fz, 'br_fz': br_fz, 'mg_fz': mg_fz, 'br_norm': br_norm_groups, 'mg_norm': mg_norm_groups,
           'bank_var': bank_var, 'diff_pairs': [(list(k), n) for k, n in diff_pairs.most_common()],
           'gubun': C(d['_gubun'] for d in erp).most_common(), 'kind': C(d['_kind'] for d in erp).most_common(),
           'channel': C(d['_channel'] for d in erp).most_common(), 'status': C(d['_status'] for d in erp).most_common()},
          open(OUT + 'summary.json', 'w'), ensure_ascii=False, default=str)
