# ERP 재적재 트랙 — 2026년 월별손익현황 작성용 워크북 (2026-10-27 사용자 요청).
#   1 홈택스 기준 매출·매입(ERP 참고)  2 통장·법인카드 분개 입력 시트(전 건)  3 대출관리현황  4 수식으로 취합되는 월별손익현황
# 환경변수 SB_URL·SB_KEY 필요. 사용: set -a && source sb.env && set +a && python3 build_pl_workbook_2026.py <출력폴더>
# 이후 반드시 recalc(LibreOffice)로 수식 계산 후 전달. 계정 배정은 사용자 확정값만 미리 채우고 나머지는 빈칸(노란색).
import os, sys, urllib.request, urllib.parse, json, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

U = os.environ['SB_URL']; K = os.environ['SB_KEY']
import argparse, calendar
ap = argparse.ArgumentParser(description='2026 월별손익현황 작성용 워크북 생성')
ap.add_argument('out'); ap.add_argument('--ledger', help='세무사 거래처별 잔액 현황 ledger.json'); ap.add_argument('--fs', help='fs_2025.json')
ap.add_argument('--prev', help='이전에 전달한 워크북(사용자 입력값을 이어받음)'); ap.add_argument('--to', default='2026-09', help='마지막 월 YYYY-MM')
args = ap.parse_args()
OUT = args.out.rstrip('/') + '/'
TO_M = int(args.to[5:7]); FROM = '2026-01-01'; TO = f'2026-{TO_M:02d}-{calendar.monthrange(2026, TO_M)[1]:02d}'
MONTHS = [f'2026-{m:02d}' for m in range(1, TO_M + 1)]
NM = len(MONTHS); TOT = 2 + NM; LASTM = L(1 + NM)   # 합계 열 번호, 마지막 월 열 글자
_ERP_S = [377912660, 686604750, 237037200, 199354220, 203727910, 219279420, 189521600, 191740000, 945819300]  # 사용자 2026 ERP 파일(부가세 포함, 1~9월)
_ERP_C = [305111041, 549841718, 198957523, 166550663, 172596581, 185525370, 155755396, 161082287, 766477159]
ERP_FILE_SALES = (_ERP_S + [0] * 12)[:NM]; ERP_FILE_COGS = (_ERP_C + [0] * 12)[:NM]   # 10월 이후는 파일 없음 → 0
TRANSFER = '0000 내부이체(손익 제외)'

def fetch_all(path):
    out = []; start = 0
    while True:
        url = U + '/rest/v1/' + urllib.parse.quote(path + '&order=id', safe="=&?.,()*:/'%!")
        r = urllib.request.Request(url, headers={'apikey': K, 'Authorization': 'Bearer ' + K, 'Range': f'{start}-{start + 999}'})
        with urllib.request.urlopen(r) as res: rows = json.loads(res.read())
        out += rows
        if len(rows) < 1000: break
        start += 1000
    return out
num = lambda v: float(v or 0)
M = lambda d: (d or '')[:7]

accounts = sorted(fetch_all('accounts?select=id,code,name,type,is_active'), key=lambda a: a['code'])
acc_by_id = {a['id']: a for a in accounts}
def acc_label(i): return f"{acc_by_id[i]['code']} {acc_by_id[i]['name']}" if i and i in acc_by_id else ''
def pl_group(a):
    c = a['code']
    if a['type'] in ('asset', 'liability', 'equity'): return '손익 제외(재무상태)'
    if c in ('4001', '4004'): return '매출(홈택스 대체·제외)'
    if c in ('5001', '5002'): return '매출원가'
    if a['type'] == 'income': return '영업외수익'
    if c.startswith('53'): return '영업외비용'
    return '판매관리비'
banks = {b['id']: b for b in fetch_all('bank_accounts?select=id,bank_name,account_type,account_number')}
cards = {c['id']: c for c in fetch_all('card_accounts?select=id,card_company,card_number,alias')}
bank_label = lambda i: (lambda b: f"{b['bank_name']} {'마이너스' if b['account_type'] == 'overdraft' else '보통'} …{(b['account_number'] or '')[-4:]}")(banks[i]) if i in banks else ''
card_label = lambda i: (lambda c: f"{c['card_company']} …{(c['card_number'] or '')[-4:]}")(cards[i]) if i in cards else ''
STATUS = {'confirmed': '확정', 'reviewed': '검토', 'pending': '대기'}

txs = fetch_all(f'transactions?select=id,tx_date,tx_time,bank_account_id,description,counterparty_name,amount_in,amount_out,tx_kind,transfer_pair_id,status,confirmed_account_id,suggested_account_id,memo&tx_date=gte.{FROM}&tx_date=lte.{TO}')
txs.sort(key=lambda t: (t['tx_date'], t['tx_time'] or '', t['id']))
ces = fetch_all(f'card_expenses?select=id,tx_date,tx_time,card_account_id,merchant_name,merchant_category,user_name,approved_amount,cancel_amount,tax_amount,classify_status,confirmed_account_id,suggested_account_id,memo&tx_date=gte.{FROM}&tx_date=lte.{TO}')
ces.sort(key=lambda c: (c['tx_date'], c['tx_time'] or '', c['id']))
tis = fetch_all(f'tax_invoices?select=id,issue_date,direction,counterparty_name,counterparty_biz_number,item_name,tax_type,supply_amount,tax_amount,total_amount,confirmed_account_id,payment_status&issue_date=gte.{FROM}&issue_date=lte.{TO}')
tis.sort(key=lambda t: (t['issue_date'], t['id']))
purchases = [t for t in tis if t['direction'] != 'sales']
sales_inv = [t for t in tis if t['direction'] == 'sales']
css = fetch_all(f'card_sales?select=tx_date,amount,supply_amount,tax_amount,transaction_type,cancelled_at&tx_date=gte.{FROM}&tx_date=lte.{TO}')
crs = fetch_all(f'cash_receipts?select=tx_date,direction,amount,supply_amount,tax_amount&tx_date=gte.{FROM}&tx_date=lte.{TO}')
loans = sorted(fetch_all('loans?select=*'), key=lambda l: l['seq'] or 0)
# 계산서 대금으로 매칭된 통장 거래 (tax_invoice_payments + tax_invoices.matched_transaction_id) — 손익에서 제외(계산서가 비용·원가를 담당)
matched_tx = set(p['transaction_id'] for p in fetch_all('tax_invoice_payments?select=transaction_id') if p['transaction_id'])
matched_tx |= set(t['matched_transaction_id'] for t in fetch_all('tax_invoices?select=matched_transaction_id&matched_transaction_id=not.is.null') if t['matched_transaction_id'])
# 추정 분류 힌트용 거래처명 (계산서 2025-10 이후)
import re
_norm = lambda x: re.sub(r'\s|\(주\)|주식회사|㈜|\(유\)', '', x or '')
inv_names = fetch_all('tax_invoices?select=direction,counterparty_name&issue_date=gte.2025-10-01')
PUR_NAMES = set(_norm(t['counterparty_name']) for t in inv_names if t['direction'] != 'sales') - {''}
SALE_NAMES = set(_norm(t['counterparty_name']) for t in inv_names if t['direction'] == 'sales') - {''}
acc_by_code = {a['code']: f"{a['code']} {a['name']}" for a in accounts}
# 세무사 2025 거래처별 잔액 현황(ledger.json, 선택 인자) → 거래처→계정 규칙 (scripts/erp-reload/ledger_vendor_balance_2025.py 산출)
LEDGER_MAP = {'상품': '5002', '지급임차료': '5109', '전력비': '5203', '세금과공과금': '5111', '접대비(기업업무추진비)': '5107', '렌트료': '5114', '이자비용': '5301', '잡손실': '5303'}
ledger_rule = {}
if args.ledger and os.path.exists(args.ledger):
    _rows = json.load(open(args.ledger))['rows']
    _tmp = {}
    for x in _rows:
        c = int(x['acode'])
        if c == 14600 or 80000 <= c < 90000 or c in (93100, 95000, 96000):
            _tmp.setdefault(_norm(x['vendor']), []).append((x['aname'], x['debit']))
    for k, v in _tmp.items():
        v.sort(key=lambda t: -t[1]); name = v[0][0]
        code = LEDGER_MAP.get(name) or next((a['code'] for a in accounts if a['name'] == name), None)
        if code and k: ledger_rule[k] = (acc_by_code[code], name)
def ledger_hint(name):
    n = _norm(name)
    if not n or not ledger_rule: return None
    if n in ledger_rule: return ledger_rule[n]
    if len(n) >= 3:
        for k, v in ledger_rule.items():
            if len(k) >= 3 and (n in k or k in n): return v
    return None
print('세무사 규칙 거래처', len(ledger_rule))
FS = json.load(open(args.fs)) if args.fs and os.path.exists(args.fs) else None
FS_ROW = {}
def _name_hit(n, names): return bool(n) and len(n) >= 3 and any(n == q or n in q or q in n for q in names)
def hint_for(t):
    """미확정 통장 거래의 추정 계정(선택값)과 근거. 확정 건은 빈칸. 시스템은 추천만, 확정은 사용자."""
    if t['confirmed_account_id']: return '', ''
    is_in = num(t['amount_in']) > 0
    if t['transfer_pair_id']: return TRANSFER, '이체쌍(시스템 판정)'
    if t['id'] in matched_tx: return (acc_by_code['1101'] if is_in else acc_by_code['2001']), '계산서 매칭 건 — 손익 제외(계정은 참고)'
    text = f"{t['counterparty_name'] or ''} {t['description'] or ''} {t['tx_kind'] or ''}"
    n = _norm(t['counterparty_name'])
    if '다올' in text: return TRANSFER, '자사 계좌 간 이체 추정(이체쌍 미감지) — 상대 계좌 확인'
    if re.search(r'카드', text) and not is_in: return acc_by_code['2001'], '카드대금 결제 추정(카드 시트에서 비용 반영) — 손익 제외'
    if re.search(r'대출이자|이자', text) and not is_in: return acc_by_code['5301'], '대출 이자 추정'
    if re.search(r'대출|원리금|중진공|보증|상환', text) and not is_in: return acc_by_code['2004'], '대출 원리금 상환 추정 — 이자 포함이면 5301 분리'
    if re.search(r'부가세|부가가치세', text): return acc_by_code['2003'], '부가세 납부 추정 — 손익 제외'
    if re.search(r'원천|소득세|지방소득세|주민세', text): return acc_by_code['2001'], '원천세 납부 추정(예수금 성격) — 손익 제외'
    if re.search(r'급여|월급|상여|퇴직', text) and not is_in: return acc_by_code['5101'], '급여 추정'
    if re.search(r'건강보험|국민연금|고용보험|산재|4대|근로복지', text) and not is_in: return acc_by_code['5102'], '4대보험 추정 — 직원 부담분은 예수금(손익 제외)로 분리 필요'
    if re.search(r'세금|지방세|국세|공과|자동차세|재산세|등록면허', text) and not is_in: return acc_by_code['5111'], '세금과공과 추정'
    if not is_in and _name_hit(n, PUR_NAMES): return acc_by_code['2001'], '매입 계산서 거래처 대금 추정 — 손익 제외(계산서가 비용 담당)'
    if is_in and _name_hit(n, SALE_NAMES): return acc_by_code['1101'], '매출 계산서 거래처 입금 추정 — 손익 제외(매출은 홈택스)'
    lh = ledger_hint(t['counterparty_name']) if not is_in else None
    if lh: return lh[0], f'세무사 2025 규칙(거래처 주계정 {lh[1]}) — 계산서 거래처면 2001'
    if t['suggested_account_id'] and t['suggested_account_id'] in acc_by_id: return acc_label(t['suggested_account_id']), '시스템 키워드 추천'
    if is_in: return '', '힌트 없음 — 매출 입금이면 1101 외상매출금(손익 제외)'
    return '', '힌트 없음 — 증빙 없는 지출이면 비용 계정, 거래처 대금이면 2001'
print('tx', len(txs), 'card', len(ces), 'purchase inv', len(purchases), 'sales inv', len(sales_inv), 'card sales', len(css), 'cash', len(crs), 'loans', len(loans))

# ── 이전 워크북의 사용자 입력 이어받기 (ID 열이 있으면 ID로, 없으면 복합 키로) ──
PREV = {'tx': {}, 'cd': {}, 'pi': {}, 'loan_master': {}, 'loan_grid': {}}
def _k(*xs):
    out = []
    for x in xs:
        if isinstance(x, float) and x.is_integer(): x = int(x)
        out.append('' if x is None else str(x).strip())
    return tuple(out)
def _amt(v): return int(v) if isinstance(v, (int, float)) and float(v).is_integer() else (round(float(v), 2) if isinstance(v, (int, float)) else '')
if args.prev and os.path.exists(args.prev):
    pw = load_workbook(args.prev, data_only=False)
    def grab(sheet, hdr_row, key_cols, val_col, id_col_name='ID'):
        ws = pw[sheet]; hdrs = [c.value for c in ws[hdr_row]]
        idc = hdrs.index(id_col_name) if id_col_name in hdrs else None
        by_id, by_key = {}, collections.defaultdict(list)
        for row in ws.iter_rows(min_row=hdr_row + 1, values_only=True):
            v = row[val_col]
            if v in (None, ''): continue
            if idc is not None and row[idc]: by_id[str(row[idc])] = str(v).strip()
            else: by_key[_k(*[row[c] for c in key_cols])].append(str(v).strip())
        return by_id, by_key
    PREV['tx'] = grab('통장 분개입력', 2, [1, 2, 3, 4, 5, 6, 7], 13)
    PREV['cd'] = grab('카드 분개입력', 2, [1, 2, 3, 4, 7, 8], 14)
    PREV['pi'] = grab('매입계산서 분류', 2, [1, 2, 3, 4, 6, 8], 11)
    wl = pw['대출관리현황']
    for row in wl.iter_rows(min_row=4, values_only=True):
        if row[1] and isinstance(row[0], (int, float)):
            PREV['loan_master'][str(row[1]).strip()] = {j: row[j - 1] for j in (8, 10, 11, 12) if row[j - 1] not in (None, '') and not str(row[j - 1]).startswith('=')}
        if row[0] and isinstance(row[0], str) and re.match(r'^\d+\. ', row[0]):
            vals = {}
            for j in range(2, 2 + 12):
                v = row[j - 1] if j - 1 < len(row) else None
                if isinstance(v, (int, float)): vals[j] = v
            if vals: PREV['loan_grid'].setdefault(row[0].strip(), []).append(vals)
    print('이전 워크북 입력값:', {k: (len(v[0]) + sum(len(x) for x in v[1].values())) if isinstance(v, tuple) else len(v) for k, v in PREV.items()})
def prev_val(kind, rid, key):
    by_id, by_key = PREV[kind] if PREV[kind] else ({}, {})
    if rid in by_id: return by_id[rid]
    lst = by_key.get(key)
    return lst.pop(0) if lst else ''

# ── 스타일 ──
FONT = '맑은 고딕'
f = lambda **kw: Font(name=FONT, **{'size': 10, **kw})
B = f(bold=True); TITLE = f(bold=True, size=13); NOTE = f(italic=True, color='555555'); GREEN = f(color='008000'); BLUE = f(color='0000FF')
HFILL = PatternFill('solid', fgColor='DDE3EA'); YELLOW = PatternFill('solid', fgColor='FFFF99'); SUBFILL = PatternFill('solid', fgColor='F3F4F6')
NUM = '#,##0;(#,##0);-'; PCT = '0.0%'
thin = Side(style='thin', color='BBBBBB'); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
def hdr(ws, r, labels, widths=None):
    for j, h in enumerate(labels, 1):
        c = ws.cell(r, j, h); c.font = B; c.fill = HFILL; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = BOX
    if widths:
        for j, w in enumerate(widths, 1): ws.column_dimensions[L(j)].width = w
def month_hdr(ws, r, first='항목'):
    hdr(ws, r, [first] + [f'{int(m[5:])}월' for m in MONTHS] + ['합계'])
    ws.column_dimensions['A'].width = 46
    for j in range(2, TOT + 1): ws.column_dimensions[L(j)].width = 14

wb = Workbook()
ws_guide = wb.active; ws_guide.title = '안내'
ws_pl = wb.create_sheet('월별손익현황')
ws_ht = wb.create_sheet('홈택스 매출매입')
ws_tx = wb.create_sheet('통장 분개입력')
ws_cd = wb.create_sheet('카드 분개입력')
ws_pi = wb.create_sheet('매입계산서 분류')
ws_ln = wb.create_sheet('대출관리현황')
ws_fs = wb.create_sheet('2025 재무제표(세무사)') if FS else None
ws_ac = wb.create_sheet('계정과목')

# ── 계정과목 ──
hdr(ws_ac, 1, ['코드', '계정명', '유형', '손익 구분', '선택값(입력 시트에서 고르는 값)'], [8, 18, 10, 22, 34])
ws_ac.append(['0000', '내부이체(손익 제외)', '-', '손익 제외(이체)', TRANSFER])
for a in accounts:
    ws_ac.append([a['code'], a['name'], a['type'], pl_group(a), f"{a['code']} {a['name']}"])
AC_LAST = ws_ac.max_row
for row in ws_ac.iter_rows(min_row=2, max_row=AC_LAST):
    for c in row: c.font = f()
ws_ac.freeze_panes = 'A2'
ws_ac.cell(AC_LAST + 2, 1, '손익 구분 규칙: 자산·부채·자본 계정과 내부이체는 손익 제외 / 4001·4004 매출은 홈택스 매출로 대체하므로 제외 / 5001·5002 = 매출원가 / 51xx·52xx = 판매관리비 / 53xx = 영업외비용 / 그 외 수익 = 영업외수익. (시스템 월별 손익현황 lib/pl-report.ts와 같은 규칙)').font = NOTE
AC_RANGE = f"계정과목!$E$2:$E${AC_LAST}"; AC_GROUP = f"계정과목!$D$2:$D${AC_LAST}"
dv_formula = f"={AC_RANGE}"

def group_formula(cell): return f'=IF({cell}="","미분류",IFERROR(INDEX({AC_GROUP},MATCH({cell},{AC_RANGE},0)),"목록에 없음"))'

# ── 통장 분개입력 ──
TX_H = ['월', '거래일', '시간', '계좌', '적요', '거래처', '입금', '출금', '거래구분', '이체쌍', '현재 상태', '기존 확정 계정', '시스템 추천 계정', '계정과목(입력)', '손익 구분(자동)', '메모', '계산서 매칭', '추정 계정(힌트)', '추정 근거', 'ID']
ws_tx.cell(1, 1, f'통장 입출금 전 건 ({FROM} ~ {TO}). N열 "계정과목(입력)"을 계정과목 시트의 선택값으로 채운다. 노란 칸 = 미확정(입력 필요). 기존 확정값은 미리 채워져 있고 수정 가능. 이체쌍은 내부이체로 미리 채움. Q열 "계산서 매칭"=Y인 건은 계산서 대금 입출금이라 계정과 무관하게 손익에서 제외된다(계산서 분류가 비용·원가를 담당 — 이중 반영 방지). R·S열 "추정 계정(힌트)·근거"는 규칙 기반 추정이다 — 맞으면 N열에 옮기고 아니면 고친다(추정은 확정이 아님).').font = NOTE
hdr(ws_tx, 2, TX_H, [9, 11, 9, 20, 28, 18, 14, 14, 10, 7, 9, 20, 20, 26, 20, 20, 9, 26, 44, 10])
r = 3
for t in txs:
    conf = acc_label(t['confirmed_account_id']); sug = acc_label(t['suggested_account_id'])
    pre = conf or (TRANSFER if t['transfer_pair_id'] else '')
    cells = [M(t['tx_date']), t['tx_date'], (t['tx_time'] or '')[:8], bank_label(t['bank_account_id']), t['description'], t['counterparty_name'],
             num(t['amount_in']) or None, num(t['amount_out']) or None]
    pv = prev_val('tx', t['id'], _k(*cells[1:6], _amt(cells[6]) if cells[6] else '', _amt(cells[7]) if cells[7] else ''))
    if pv: pre = pv
    ws_tx.append(cells + [t['tx_kind'], 'Y' if t['transfer_pair_id'] else '', STATUS.get(t['status'], t['status']),
                  conf, sug, pre, f'=IF(Q{r}="Y","손익 제외(계산서 매칭)",{group_formula(f"N{r}")[1:]})', t['memo'], 'Y' if t['id'] in matched_tx else '', *hint_for(t), t['id']])
    if not pre: ws_tx.cell(r, 14).fill = YELLOW
    r += 1
TX_LAST = r - 1
for row in ws_tx.iter_rows(min_row=3, max_row=TX_LAST):
    for c in row: c.font = f()
    row[6].number_format = NUM; row[7].number_format = NUM
dv_tx = DataValidation(type='list', formula1=dv_formula, allow_blank=True, showErrorMessage=True, errorTitle='계정과목', error='계정과목 시트의 선택값 중에서 고르세요.')
ws_tx.add_data_validation(dv_tx); dv_tx.add(f'N3:N{TX_LAST}')
ws_tx.freeze_panes = 'E3'; ws_tx.auto_filter.ref = f'A2:S{TX_LAST}'; ws_tx.column_dimensions['T'].hidden = True

# ── 카드 분개입력 ──
CD_H = ['월', '거래일', '시간', '카드', '가맹점', '업종', '사용자', '승인금액', '취소금액', '순액(자동)', '세액', '현재 상태', '기존 확정 계정', '시스템 추천 계정', '계정과목(입력)', '손익 구분(자동)', '메모', 'ID']
ws_cd.cell(1, 1, f'법인카드 사용내역 전 건 ({FROM} ~ {TO}). O열 "계정과목(입력)"을 채운다. 노란 칸 = 미확정. 순액 = 승인금액 - 취소금액(수식). 카드대금 결제 출금은 통장 시트에서 2001 미지급금 등 손익 제외 계정으로 처리해 이중 반영을 막는다.').font = NOTE
hdr(ws_cd, 2, CD_H, [9, 11, 9, 16, 24, 14, 10, 13, 12, 13, 11, 9, 20, 20, 26, 18, 18, 10])
r = 3
for c in ces:
    conf = acc_label(c['confirmed_account_id']); sug = acc_label(c['suggested_account_id'])
    head = [M(c['tx_date']), c['tx_date'], (c['tx_time'] or '')[:8], card_label(c['card_account_id']), c['merchant_name'], c['merchant_category'], c['user_name'], num(c['approved_amount']), num(c['cancel_amount'])]
    pv = prev_val('cd', c['id'], _k(head[1], head[2], head[3], head[4], _amt(head[7]), _amt(head[8])))
    val = pv or conf
    ws_cd.append(head + [f'=H{r}-I{r}', num(c['tax_amount']), STATUS.get(c['classify_status'], c['classify_status']), conf, sug, val, group_formula(f'O{r}'), c['memo'], c['id']])
    if not val: ws_cd.cell(r, 15).fill = YELLOW
    r += 1
CD_LAST = r - 1
for row in ws_cd.iter_rows(min_row=3, max_row=CD_LAST):
    for c in row: c.font = f()
    for j in (7, 8, 9, 10): row[j].number_format = NUM
dv_cd = DataValidation(type='list', formula1=dv_formula, allow_blank=True, showErrorMessage=True, errorTitle='계정과목', error='계정과목 시트의 선택값 중에서 고르세요.')
ws_cd.add_data_validation(dv_cd); dv_cd.add(f'O3:O{CD_LAST}')
ws_cd.freeze_panes = 'E3'; ws_cd.auto_filter.ref = f'A2:Q{CD_LAST}'; ws_cd.column_dimensions['R'].hidden = True

# ── 매입계산서 분류 ──
PI_H = ['월', '발행일', '거래처', '사업자번호', '품목', '과세/면세', '공급가액', '세액', '합계', '결제 매칭', '기존 확정 계정', '계정과목(입력)', '손익 구분(자동)', '추정 계정(힌트)', '추정 근거', 'ID']
ws_pi.cell(1, 1, f'매입 세금계산서 전 건 (발행일 {FROM} ~ {TO}). L열에 계정과목을 채운다(N·O열 세무사 2025 규칙 힌트 참고) — 상품 매입은 5002 상품매입(매출원가), 임차료·용역 등은 해당 비용 계정. 노란 칸 = 미분류. 계산서 대금을 통장에서 지급한 출금은 통장 시트에서 2001 미지급금(손익 제외)으로 처리한다.').font = NOTE
hdr(ws_pi, 2, PI_H, [9, 11, 24, 14, 24, 9, 14, 12, 14, 10, 20, 26, 18, 26, 40, 10])
r = 3
for t in purchases:
    conf = acc_label(t['confirmed_account_id'])
    pv = prev_val('pi', t['id'], _k(t['issue_date'], t['counterparty_name'], t['counterparty_biz_number'], t['item_name'], _amt(num(t['supply_amount'])), _amt(num(t['total_amount']))))
    val = pv or conf
    ws_pi.append([M(t['issue_date']), t['issue_date'], t['counterparty_name'], t['counterparty_biz_number'], t['item_name'], '과세' if t['tax_type'] == 'taxable' else '면세',
                  num(t['supply_amount']), num(t['tax_amount']), num(t['total_amount']), '매칭' if t['payment_status'] == 'matched' else '미매칭', conf, val, group_formula(f'L{r}'),
                  *(('', '') if val else ((lambda lh: (lh[0], f'세무사 2025 규칙(거래처 주계정 {lh[1]})') if lh else ('', '규칙 없음 — 품목으로 판단(상품이면 5002)'))(ledger_hint(t['counterparty_name'])))), t['id']])
    if not val: ws_pi.cell(r, 12).fill = YELLOW
    r += 1
PI_LAST = r - 1
for row in ws_pi.iter_rows(min_row=3, max_row=PI_LAST):
    for c in row: c.font = f()
    for j in (6, 7, 8): row[j].number_format = NUM
dv_pi = DataValidation(type='list', formula1=dv_formula, allow_blank=True, showErrorMessage=True, errorTitle='계정과목', error='계정과목 시트의 선택값 중에서 고르세요.')
ws_pi.add_data_validation(dv_pi); dv_pi.add(f'L3:L{PI_LAST}')
ws_pi.freeze_panes = 'C3'; ws_pi.auto_filter.ref = f'A2:O{PI_LAST}'; ws_pi.column_dimensions['P'].hidden = True

# 범위 문자열
TXm = f"'통장 분개입력'!$A$3:$A${TX_LAST}"; TXin = f"'통장 분개입력'!$G$3:$G${TX_LAST}"; TXout = f"'통장 분개입력'!$H$3:$H${TX_LAST}"; TXacc = f"'통장 분개입력'!$N$3:$N${TX_LAST}"; TXgrp = f"'통장 분개입력'!$O$3:$O${TX_LAST}"; TXpair = f"'통장 분개입력'!$J$3:$J${TX_LAST}"; TXmatch = f"'통장 분개입력'!$Q$3:$Q${TX_LAST}"
CDm = f"'카드 분개입력'!$A$3:$A${CD_LAST}"; CDnet = f"'카드 분개입력'!$J$3:$J${CD_LAST}"; CDacc = f"'카드 분개입력'!$O$3:$O${CD_LAST}"; CDgrp = f"'카드 분개입력'!$P$3:$P${CD_LAST}"; CDtax = f"'카드 분개입력'!$K$3:$K${CD_LAST}"
PIm = f"'매입계산서 분류'!$A$3:$A${PI_LAST}"; PIsup = f"'매입계산서 분류'!$G$3:$G${PI_LAST}"; PItax = f"'매입계산서 분류'!$H$3:$H${PI_LAST}"; PIacc = f"'매입계산서 분류'!$L$3:$L${PI_LAST}"; PIgrp = f"'매입계산서 분류'!$M$3:$M${PI_LAST}"; PIkind = f"'매입계산서 분류'!$F$3:$F${PI_LAST}"

# ── 홈택스 매출매입 ──
ws_ht.cell(1, 1, '홈택스 기준 매출·매입 (2026-01 ~ 09). 매출 = 매출 세금계산서(발행일) + 카드 매출(승인-취소) + 현금영수증(39건 전량). 매입 = 매입 세금계산서 + 법인카드 + 현금영수증. 추출 2026-10-27, 매출 쪽 숫자는 시스템 원본 집계값, 매입 쪽은 입력 시트에서 수식으로 집계.').font = NOTE
ws_ht.cell(2, 1, 'ERP 기준 매출·매입원가는 참고용(사용자 제공 2026 ERP 판매내역 파일, 부가세 포함). 손익현황은 공급가액(세액 제외) 기준으로 가져간다.').font = NOTE
month_hdr(ws_ht, 4)
S = collections.defaultdict(lambda: {m: 0.0 for m in MONTHS})
for t in sales_inv:
    m = M(t['issue_date']); kind = '과세' if t['tax_type'] == 'taxable' else '면세'
    S[f'계산서 {kind} 공급가액'][m] += num(t['supply_amount']); S[f'계산서 {kind} 세액'][m] += num(t['tax_amount']); S[f'계산서 {kind} 건수'][m] += 1
for c in css:
    m = M(c['tx_date']); sign = -1 if c['transaction_type'] == 'cancel' else 1
    if c['cancelled_at'] and sign == 1: continue
    if num(c['tax_amount']) > 0: S['카드 매출 과세 공급가액'][m] += sign * num(c['supply_amount']); S['카드 매출 세액'][m] += sign * num(c['tax_amount'])
    else: S['카드 매출 세액 0 금액(면세 또는 미분리)'][m] += sign * num(c['amount'])
    S['카드 매출 건수(승인)'][m] += 1 if sign == 1 else 0
for c in crs:
    m = M(c['tx_date']); D = '매출' if c['direction'] == 'sales' else '매입'
    S[f'현금영수증 {D} 공급가액'][m] += num(c['supply_amount']); S[f'현금영수증 {D} 세액'][m] += num(c['tax_amount']); S[f'현금영수증 {D} 건수'][m] += 1
HT = {}  # 라벨 → 행번호
r = 5
def ht_row(label, values=None, formula=None, bold=False, fmt=NUM, total=True):
    global r
    ws_ht.cell(r, 1, label).font = f(bold=bold)
    for j, m in enumerate(MONTHS, 2):
        col = L(j)
        c = ws_ht.cell(r, j, formula(col, m) if formula else round(values[m]) if values else 0)
        c.number_format = fmt; c.font = f(bold=bold) if not (formula and "'" in formula(col, m)) else GREEN
        if bold: c.font = f(bold=True)
    if total:
        c = ws_ht.cell(r, TOT, f'=SUM(B{r}:{LASTM}{r})'); c.number_format = fmt; c.font = f(bold=True)
    HT[label] = r; r += 1
def section(title):
    global r
    c = ws_ht.cell(r, 1, title); c.font = B; c.fill = SUBFILL
    for j in range(2, TOT + 1): ws_ht.cell(r, j).fill = SUBFILL
    r += 1
section('매출 (홈택스 기준)')
ht_row('계산서 과세 공급가액', S['계산서 과세 공급가액']); ht_row('계산서 과세 세액', S['계산서 과세 세액']); ht_row('계산서 과세 건수', S['계산서 과세 건수'], fmt='#,##0')
ht_row('계산서 면세 공급가액', S['계산서 면세 공급가액']); ht_row('계산서 면세 건수', S['계산서 면세 건수'], fmt='#,##0')
ht_row('카드 매출 과세 공급가액', S['카드 매출 과세 공급가액']); ht_row('카드 매출 세액', S['카드 매출 세액']); ht_row('카드 매출 세액 0 금액(면세 또는 미분리)', S['카드 매출 세액 0 금액(면세 또는 미분리)']); ht_row('카드 매출 건수(승인)', S['카드 매출 건수(승인)'], fmt='#,##0')
ht_row('현금영수증 매출 공급가액', S['현금영수증 매출 공급가액']); ht_row('현금영수증 매출 세액', S['현금영수증 매출 세액']); ht_row('현금영수증 매출 건수', S['현금영수증 매출 건수'], fmt='#,##0')
ht_row('매출 공급가액 합계(세액 제외)', formula=lambda col, m: f"={col}{HT['계산서 과세 공급가액']}+{col}{HT['계산서 면세 공급가액']}+{col}{HT['카드 매출 과세 공급가액']}+{col}{HT['카드 매출 세액 0 금액(면세 또는 미분리)']}+{col}{HT['현금영수증 매출 공급가액']}", bold=True)
ht_row('매출 세액 합계', formula=lambda col, m: f"={col}{HT['계산서 과세 세액']}+{col}{HT['카드 매출 세액']}+{col}{HT['현금영수증 매출 세액']}")
ht_row('매출 합계(부가세 포함)', formula=lambda col, m: f"={col}{HT['매출 공급가액 합계(세액 제외)']}+{col}{HT['매출 세액 합계']}", bold=True)
ht_row('참고: ERP 매출(판매내역 파일, 부가세 포함; 10월 이후 파일 없음=0)', {m: ERP_FILE_SALES[i] for i, m in enumerate(MONTHS)})
ht_row('참고: 홈택스 매출 / ERP 매출', formula=lambda col, m: f"=IF({col}{HT['참고: ERP 매출(판매내역 파일, 부가세 포함; 10월 이후 파일 없음=0)']}=0,0,{col}{HT['매출 합계(부가세 포함)']}/{col}{HT['참고: ERP 매출(판매내역 파일, 부가세 포함; 10월 이후 파일 없음=0)']})", fmt=PCT, total=False)
ws_ht.cell(HT['참고: 홈택스 매출 / ERP 매출'], TOT, f"=IF({L(TOT)}{HT['참고: ERP 매출(판매내역 파일, 부가세 포함; 10월 이후 파일 없음=0)']}=0,0,{L(TOT)}{HT['매출 합계(부가세 포함)']}/{L(TOT)}{HT['참고: ERP 매출(판매내역 파일, 부가세 포함; 10월 이후 파일 없음=0)']})").number_format = PCT
r += 1
section('매입 (홈택스 기준) — 입력 시트에서 집계')
ht_row('매입 계산서 과세 공급가액', formula=lambda col, m: f'=SUMIFS({PIsup},{PIm},"{m}",{PIkind},"과세")')
ht_row('매입 계산서 과세 세액', formula=lambda col, m: f'=SUMIFS({PItax},{PIm},"{m}",{PIkind},"과세")')
ht_row('매입 계산서 면세 공급가액', formula=lambda col, m: f'=SUMIFS({PIsup},{PIm},"{m}",{PIkind},"면세")')
ht_row('매입 계산서 건수', formula=lambda col, m: f'=COUNTIFS({PIm},"{m}")', fmt='#,##0')
ht_row('법인카드 사용액(승인-취소)', formula=lambda col, m: f'=SUMIFS({CDnet},{CDm},"{m}")')
ht_row('법인카드 세액(카드사 표기)', formula=lambda col, m: f'=SUMIFS({CDtax},{CDm},"{m}")')
ht_row('법인카드 건수', formula=lambda col, m: f'=COUNTIFS({CDm},"{m}")', fmt='#,##0')
ht_row('현금영수증 매입 공급가액', S['현금영수증 매입 공급가액']); ht_row('현금영수증 매입 세액', S['현금영수증 매입 세액'])
ht_row('매입 공급가액 합계(세액 제외)', formula=lambda col, m: f"={col}{HT['매입 계산서 과세 공급가액']}+{col}{HT['매입 계산서 면세 공급가액']}+{col}{HT['법인카드 사용액(승인-취소)']}-{col}{HT['법인카드 세액(카드사 표기)']}+{col}{HT['현금영수증 매입 공급가액']}", bold=True)
ht_row('매입 합계(부가세 포함)', formula=lambda col, m: f"={col}{HT['매입 계산서 과세 공급가액']}+{col}{HT['매입 계산서 과세 세액']}+{col}{HT['매입 계산서 면세 공급가액']}+{col}{HT['법인카드 사용액(승인-취소)']}+{col}{HT['현금영수증 매입 공급가액']}+{col}{HT['현금영수증 매입 세액']}", bold=True)
ht_row('참고: ERP 매입원가(판매내역 파일)', {m: ERP_FILE_COGS[i] for i, m in enumerate(MONTHS)})
r += 1
section('참고: 통장 (내부 이체쌍 제외)')
ht_row('통장 입금 합계', formula=lambda col, m: f'=SUMIFS({TXin},{TXm},"{m}",{TXpair},"<>Y")')
ht_row('통장 출금 합계', formula=lambda col, m: f'=SUMIFS({TXout},{TXm},"{m}",{TXpair},"<>Y")')
ws_ht.cell(r + 1, 1, '법인카드 사용액은 홈택스 매입 합계에 포함되지만, 손익현황에서는 카드 분개입력 시트의 계정별로 판매관리비에 반영된다(매입 합계와 이중 반영 아님).').font = NOTE
ws_ht.freeze_panes = 'B5'

# ── 대출관리현황 ──
ws_ln.cell(1, 1, '대출관리현황 — 시스템 대출 마스터(2026-08-24~26 갱신). 8월 이후 미갱신: 노란 칸(잔액·월 이자)을 갱신한다. 한도대출(마이너스)은 사용액에 따라 이자가 달라 통장 5301 이자비용 실적으로 확인.').font = NOTE
LN_H = ['순번', '대출명', '은행', '종류', '상태', '최초 금액', '한도', '현재 잔액', '잔액 기준일', '금리(%)', '월 원금상환', '월 이자(계약 기준)', '시작일', '만기일', '상환일', '마스터 갱신일', '메모']
hdr(ws_ln, 3, LN_H, [6, 30, 10, 10, 8, 14, 14, 14, 11, 8, 13, 14, 11, 11, 7, 11, 30])
r = 4; LN_FIRST = r
for l in loans:
    ws_ln.append([l['seq'], l['title'], l['bank_name'], '한도대출' if l['product_type'] == 'credit_line' else '일반대출', '정상' if l['status'] == 'active' else l['status'],
                  num(l['original_amount']) or None, num(l['credit_limit']) or None, num(l['current_balance']), l['balance_date'], l['interest_rate'], num(l['monthly_principal']), num(l['monthly_interest']),
                  l['start_date'], l['maturity_date'], l['payment_day'], (l['updated_at'] or '')[:10], l['memo']])
    for j in (6, 7, 8, 11, 12): ws_ln.cell(r, j).number_format = NUM
    for j in (8, 10, 11, 12): ws_ln.cell(r, j).fill = YELLOW
    for j, v in PREV['loan_master'].get(l['title'], {}).items(): ws_ln.cell(r, j, v)   # 사용자 갱신값 이어받기
    r += 1
LN_LAST = r - 1
ws_ln.cell(r, 2, '합계').font = B
for j in (8, 11, 12):
    c = ws_ln.cell(r, j, f'=SUM({L(j)}{LN_FIRST}:{L(j)}{LN_LAST})'); c.number_format = NUM; c.font = B
for row in ws_ln.iter_rows(min_row=4, max_row=r):
    for c in row:
        if c.font.name != FONT: c.font = f(bold=c.font.bold)
r += 2
ws_ln.cell(r, 1, '월별 이자 (계약 기준 계획, 원). 각 대출의 "월 이자(계약 기준)"를 각 월에 같은 값으로 채워 두었다 — 실제 납부액으로 고쳐 쓴다(노란 칸). 합계 행이 손익현황의 참고 행으로 간다.').font = NOTE; r += 1
month_hdr(ws_ln, r, '대출명'); ws_ln.column_dimensions['A'].width = 6; ws_ln.column_dimensions['B'].width = 30
# 월별 표는 B열에 대출명이 아니라 A열 라벨을 쓰는 month_hdr와 맞추기 위해 A열에 대출명
r += 1; LNI_FIRST = r
for i, l in enumerate(loans):
    ws_ln.cell(r, 1, f"{l['seq']}. {l['title']}").font = f()
    grid = PREV['loan_grid'].get(f"{l['seq']}. {l['title']}", [{}])
    for j in range(2, 2 + NM):
        c = ws_ln.cell(r, j, grid[0].get(j, f'=$L${LN_FIRST + i}')); c.number_format = NUM; c.fill = YELLOW; c.font = f()
    c = ws_ln.cell(r, TOT, f'=SUM(B{r}:{LASTM}{r})'); c.number_format = NUM; c.font = B
    r += 1
LNI_LAST = r - 1; LNI_TOTAL = r
ws_ln.cell(r, 1, '월 이자 계획 합계').font = B
for j in range(2, TOT + 1):
    c = ws_ln.cell(r, j, f'=SUM({L(j)}{LNI_FIRST}:{L(j)}{LNI_LAST})'); c.number_format = NUM; c.font = B
r += 2
ws_ln.cell(r, 1, '월말 잔액 (원) — 비어 있음. 은행 조회 후 입력(노란 칸). 마스터 "현재 잔액"은 2026-08 기준이라 월별로 그대로 쓰지 않는다.').font = NOTE; r += 1
month_hdr(ws_ln, r, '대출명'); ws_ln.column_dimensions['A'].width = 34; r += 1; LNB_FIRST = r
for l in loans:
    ws_ln.cell(r, 1, f"{l['seq']}. {l['title']}").font = f()
    grid = PREV['loan_grid'].get(f"{l['seq']}. {l['title']}", [{}, {}])
    bal = grid[1] if len(grid) > 1 else {}
    for j in range(2, 2 + NM):
        ws_ln.cell(r, j, bal.get(j)); ws_ln.cell(r, j).fill = YELLOW; ws_ln.cell(r, j).number_format = NUM
    r += 1
ws_ln.cell(r, 1, '월말 잔액 합계').font = B
for j in range(2, 2 + NM):
    c = ws_ln.cell(r, j, f'=SUM({L(j)}{LNB_FIRST}:{L(j)}{r - 1})'); c.number_format = NUM; c.font = B
ws_ln.freeze_panes = 'A4'

# ── 월별손익현황 ──
if FS:
    ws_fs.cell(1, 1, f"2025년 재무제표 (세무사 작성, 제11기) — 출처: {FS['source']}. 단위 원. 월별손익현황 L열의 비교 기준.").font = NOTE
    hdr(ws_fs, 3, ['손익계산서 과목', '제11기 (2025)', '제10기 (2024)', '증감'], [44, 18, 18, 16])
    rr = 4
    for name, cur, prev, kind in FS['income_statement']:
        ws_fs.cell(rr, 1, name).font = f(bold=kind == 'subtotal')
        c = ws_fs.cell(rr, 2, cur); c.number_format = NUM; c.font = f(bold=kind == 'subtotal')
        c = ws_fs.cell(rr, 3, prev); c.number_format = NUM; c.font = f(bold=kind == 'subtotal')
        c = ws_fs.cell(rr, 4, f'=B{rr}-C{rr}'); c.number_format = NUM; c.font = f()
        if kind == 'subtotal':
            for j in range(1, 5): ws_fs.cell(rr, j).fill = SUBFILL
        FS_ROW[name] = rr; rr += 1
    rr += 1; ws_fs.cell(rr, 1, FS['unmapped_note']).font = NOTE; rr += 2
    hdr(ws_fs, rr, ['재무상태표 과목', '제11기 (2025-12-31)', '제10기 (2024-12-31)', '증감']); rr += 1
    for name, cur, prev, kind in FS['balance_sheet']:
        ws_fs.cell(rr, 1, name).font = f(bold=kind in ('subtotal', 'total', 'header'))
        if cur is not None:
            c = ws_fs.cell(rr, 2, cur); c.number_format = NUM; c.font = f(bold=kind in ('subtotal', 'total'))
            c = ws_fs.cell(rr, 3, prev); c.number_format = NUM; c.font = f(bold=kind in ('subtotal', 'total'))
            c = ws_fs.cell(rr, 4, f'=B{rr}-C{rr}'); c.number_format = NUM; c.font = f()
        if kind in ('subtotal', 'total', 'header'):
            for j in range(1, 5): ws_fs.cell(rr, j).fill = SUBFILL
        rr += 1
    ws_fs.freeze_panes = 'A4'
ws_pl.cell(1, 1, f'2026년 월별손익현황 (1~{NM}월) — 홈택스 기준 매출·매입 + 통장·법인카드·매입계산서 분개 입력 시트에서 수식으로 취합. 금액 단위: 원, 공급가액(세액 제외) 기준.').font = TITLE
ws_pl.cell(2, 1, '입력 시트의 계정과목이 채워지는 대로 자동 갱신된다. 아래 "미확정 현황"이 0이 되어야 완성. 통장 입금 중 매출 입금(4001)은 홈택스 매출로 대체하므로 손익에 넣지 않는다.').font = NOTE
month_hdr(ws_pl, 4); ws_pl.column_dimensions['A'].width = 48
if FS:
    for j, h in enumerate(['2025 연간(세무사 재무제표)', '2025 월평균(÷12)', f'2026 1~{NM}월 월평균(÷{NM})'], TOT + 1):
        c = ws_pl.cell(4, j, h); c.font = B; c.fill = HFILL; c.alignment = Alignment(horizontal='center', wrap_text=True); ws_pl.column_dimensions[L(j)].width = 16
r = 5; PL = {}
def pl_row(label, formula, bold=False, fill=None, fmt=NUM, total=True, link=False):
    global r
    c = ws_pl.cell(r, 1, label); c.font = f(bold=bold)
    if fill: c.fill = fill
    for j, m in enumerate(MONTHS, 2):
        col = L(j)
        c = ws_pl.cell(r, j, formula(col, m)); c.number_format = fmt; c.font = GREEN if link else f(bold=bold)
        if fill: c.fill = fill
    if total:
        c = ws_pl.cell(r, TOT, f'=SUM(B{r}:{LASTM}{r})'); c.number_format = fmt; c.font = f(bold=True)
        if fill: c.fill = fill
        if FS and label in FS['map_to_ours']:
            refs = [f"'2025 재무제표(세무사)'!$B${FS_ROW[n]}" for n in FS['map_to_ours'][label] if n in FS_ROW]
            if refs:
                c = ws_pl.cell(r, TOT + 1, '=' + '+'.join(refs)); c.number_format = fmt; c.font = GREEN
                c = ws_pl.cell(r, TOT + 2, f'={L(TOT + 1)}{r}/12'); c.number_format = fmt; c.font = f(bold=bold)
                c = ws_pl.cell(r, TOT + 3, f'={L(TOT)}{r}/{NM}'); c.number_format = fmt; c.font = f(bold=bold)
    PL[label] = r; r += 1
def pl_section(title):
    global r
    c = ws_pl.cell(r, 1, title); c.font = B; c.fill = SUBFILL
    for j in range(2, TOT + 1): ws_pl.cell(r, j).fill = SUBFILL
    r += 1
ht = lambda label: lambda col, m: f"='홈택스 매출매입'!{col}{HT[label]}"
pl_row('매출 (홈택스 기준 공급가액)', ht('매출 공급가액 합계(세액 제외)'), bold=True, link=True)
pl_section('매출원가')
pl_row('매입 계산서 — 상품매입·매출원가 계정 분류분', lambda col, m: f'=SUMIFS({PIsup},{PIm},"{m}",{PIgrp},"매출원가")')
pl_row('매입 계산서 — 미분류분 (원가로 가정, 분류하면 자동 이동)', lambda col, m: f'=SUMIFS({PIsup},{PIm},"{m}",{PIgrp},"미분류")')
pl_row('법인카드 — 상품매입·매출원가 계정 분류분', lambda col, m: f'=SUMIFS({CDnet},{CDm},"{m}",{CDgrp},"매출원가")')
pl_row('통장 — 상품매입 계정 분류분 (계산서 미매칭 건만, 확인 필요)', lambda col, m: f'=SUMIFS({TXout},{TXm},"{m}",{TXgrp},"매출원가")-SUMIFS({TXin},{TXm},"{m}",{TXgrp},"매출원가")')
pl_row('현금영수증 매입 (공급가액)', ht('현금영수증 매입 공급가액'), link=True)
pl_row('매출원가 합계', lambda col, m: f"={col}{PL['매입 계산서 — 상품매입·매출원가 계정 분류분']}+{col}{PL['매입 계산서 — 미분류분 (원가로 가정, 분류하면 자동 이동)']}+{col}{PL['법인카드 — 상품매입·매출원가 계정 분류분']}+{col}{PL['통장 — 상품매입 계정 분류분 (계산서 미매칭 건만, 확인 필요)']}+{col}{PL['현금영수증 매입 (공급가액)']}", bold=True)
pl_row('매출총이익', lambda col, m: f"={col}{PL['매출 (홈택스 기준 공급가액)']}-{col}{PL['매출원가 합계']}", bold=True, fill=SUBFILL)
def acc_expr(label, col, m, sign_bank='out'):
    # 통장: 출금-입금(비용) 또는 입금-출금(수익) / 카드 순액 / 매입계산서 공급가액
    bank = f'SUMIFS({TXout},{TXm},"{m}",{TXacc},"{label}",{TXmatch},"<>Y")-SUMIFS({TXin},{TXm},"{m}",{TXacc},"{label}",{TXmatch},"<>Y")'
    if sign_bank == 'in': bank = f'SUMIFS({TXin},{TXm},"{m}",{TXacc},"{label}",{TXmatch},"<>Y")-SUMIFS({TXout},{TXm},"{m}",{TXacc},"{label}",{TXmatch},"<>Y")'
    return f'={bank}+SUMIFS({CDnet},{CDm},"{m}",{CDacc},"{label}")+SUMIFS({PIsup},{PIm},"{m}",{PIacc},"{label}")'
pl_section('판매관리비 (통장·카드·매입계산서 분개 — 계정과목별)')
sga = [a for a in accounts if pl_group(a) == '판매관리비']
for a in sga:
    lab = f"{a['code']} {a['name']}"; pl_row(lab, lambda col, m, lab=lab: acc_expr(lab, col, m))
SGA_FIRST = PL[f"{sga[0]['code']} {sga[0]['name']}"]; SGA_LAST = PL[f"{sga[-1]['code']} {sga[-1]['name']}"]
pl_row('판매관리비 합계', lambda col, m: f'=SUM({col}{SGA_FIRST}:{col}{SGA_LAST})', bold=True)
pl_row('영업이익', lambda col, m: f"={col}{PL['매출총이익']}-{col}{PL['판매관리비 합계']}", bold=True, fill=SUBFILL)
pl_section('영업외수익 (통장 분개)')
noi = [a for a in accounts if pl_group(a) == '영업외수익']
for a in noi:
    lab = f"{a['code']} {a['name']}"; pl_row(lab, lambda col, m, lab=lab: acc_expr(lab, col, m, 'in'))
NOI_FIRST = PL[f"{noi[0]['code']} {noi[0]['name']}"]; NOI_LAST = PL[f"{noi[-1]['code']} {noi[-1]['name']}"]
pl_row('영업외수익 합계', lambda col, m: f'=SUM({col}{NOI_FIRST}:{col}{NOI_LAST})', bold=True)
pl_section('영업외비용 (통장·카드 분개)')
noe = [a for a in accounts if pl_group(a) == '영업외비용']
for a in noe:
    lab = f"{a['code']} {a['name']}"; pl_row(lab, lambda col, m, lab=lab: acc_expr(lab, col, m))
NOE_FIRST = PL[f"{noe[0]['code']} {noe[0]['name']}"]; NOE_LAST = PL[f"{noe[-1]['code']} {noe[-1]['name']}"]
pl_row('영업외비용 합계', lambda col, m: f'=SUM({col}{NOE_FIRST}:{col}{NOE_LAST})', bold=True)
pl_row('법인세차감전순이익', lambda col, m: f"={col}{PL['영업이익']}+{col}{PL['영업외수익 합계']}-{col}{PL['영업외비용 합계']}", bold=True, fill=SUBFILL)
r += 1
pl_section('미확정 현황 (0이 되어야 완성)')
pl_row('통장 — 계정 미입력 건수', lambda col, m: f'=COUNTIFS({TXm},"{m}",{TXacc},"")', fmt='#,##0')
pl_row('통장 — 계정 미입력 출금 금액 (손익 미반영)', lambda col, m: f'=SUMIFS({TXout},{TXm},"{m}",{TXacc},"")')
pl_row('통장 — 계정 미입력 입금 금액 (손익 미반영)', lambda col, m: f'=SUMIFS({TXin},{TXm},"{m}",{TXacc},"")')
pl_row('카드 — 계정 미입력 건수', lambda col, m: f'=COUNTIFS({CDm},"{m}",{CDacc},"")', fmt='#,##0')
pl_row('카드 — 계정 미입력 순액 (손익 미반영)', lambda col, m: f'=SUMIFS({CDnet},{CDm},"{m}",{CDacc},"")')
pl_row('매입 계산서 — 계정 미입력 건수', lambda col, m: f'=COUNTIFS({PIm},"{m}",{PIacc},"")', fmt='#,##0')
pl_row('목록에 없는 계정값 건수 (오타 확인)', lambda col, m: f'=COUNTIFS({TXm},"{m}",{TXgrp},"목록에 없음")+COUNTIFS({CDm},"{m}",{CDgrp},"목록에 없음")+COUNTIFS({PIm},"{m}",{PIgrp},"목록에 없음")', fmt='#,##0')
r += 1
pl_section('참고')
pl_row('대출 이자 계획 합계 (대출관리현황, 계약 기준)', lambda col, m: f"='대출관리현황'!{col}{LNI_TOTAL}", link=True)
pl_row('5301 이자비용 실적 (분개, 위 영업외비용과 동일)', lambda col, m: f"={col}{PL['5301 이자비용']}")
pl_row('ERP 매출 (판매내역 파일, 부가세 포함, 참고)', ht('참고: ERP 매출(판매내역 파일, 부가세 포함; 10월 이후 파일 없음=0)'), link=True)
pl_row('ERP 매입원가 (판매내역 파일, 참고)', ht('참고: ERP 매입원가(판매내역 파일)'), link=True)
pl_row('홈택스 매출(부가세 포함) / ERP 매출', ht('참고: 홈택스 매출 / ERP 매출'), fmt=PCT, total=False, link=True)
pl_row('손익 제외로 분류된 통장 출금 (이체·미지급금·차입금 상환 등)', lambda col, m: f'=SUMIFS({TXout},{TXm},"{m}",{TXgrp},"손익 제외(이체)")+SUMIFS({TXout},{TXm},"{m}",{TXgrp},"손익 제외(재무상태)")')
pl_row('매출 입금으로 분류된 통장 입금 (4001·4004, 홈택스로 대체)', lambda col, m: f'=SUMIFS({TXin},{TXm},"{m}",{TXgrp},"매출(홈택스 대체·제외)")')
pl_row('계산서 매칭 통장 출금 (손익 제외 — 계산서 분류가 담당)', lambda col, m: f'=SUMIFS({TXout},{TXm},"{m}",{TXmatch},"Y")')
pl_row('  그중 비용 계정으로 확정돼 있던 금액 (이중 반영 방지로 제외됨)', lambda col, m: f'=SUMIFS({TXout},{TXm},"{m}",{TXmatch},"Y",{TXacc},"5*")')
pl_row('매입 계산서 미매칭 공급가액 (대금 출금을 2001로 분류했는지 확인)', lambda col, m: f'=SUMIFS({PIsup},{PIm},"{m}",\'매입계산서 분류\'!$J$3:$J${PI_LAST},"미매칭")')
ws_pl.freeze_panes = 'B5'

# ── 안내 ──
g = ws_guide; g.column_dimensions['A'].width = 120
lines = [
    ('2026년 월별손익현황 작성용 워크북 — 2026-10-27 추출 (daolERP 총괄)', TITLE),
    ('', None),
    ('[시트 구성]', B),
    ('월별손익현황: 결과. 모든 숫자가 수식 — 입력 시트를 채우면 자동 갱신. 아래쪽 "미확정 현황"이 전부 0이 되어야 완성.', None),
    ('홈택스 매출매입: 매출(세금계산서+카드+현금영수증)은 시스템 원본 집계값, 매입은 입력 시트에서 수식 집계. ERP 매출·원가는 참고 행.', None),
    ('통장 분개입력 / 카드 분개입력 / 매입계산서 분류: 전 건 목록. "계정과목(입력)" 열만 채운다(드롭다운, 계정과목 시트의 선택값). 노란 칸 = 미확정.', None),
    ('통장 시트 R·S열 "추정 계정(힌트)·추정 근거": 미확정 건에 규칙으로 붙인 추정(자사 이체·카드대금·대출·급여·세금·보험·계산서 거래처·세무사 2025 거래처 규칙·시스템 추천). 맞으면 N열에 복사, 아니면 고친다. 추정은 확정이 아니며 손익에는 N열 값만 쓰인다.', None),
    ('매입계산서 분류 시트 N·O열: 세무사 2025 거래처별 잔액 현황에서 뽑은 "그 거래처는 이 계정" 규칙(세무사 상품 → 5002 상품매입, 지급임차료 → 5109 등). 2025년에 거래가 없던 거래처는 규칙 없음.', None),
    ('대출관리현황: 대출 마스터(8월 갱신) + 월별 이자 계획·월말 잔액(노란 칸 갱신).', None),
    ('2025 재무제표(세무사): 손익계산서·재무상태표 전사본. 월별손익현황 L·M열(2025 연간·월평균)이 여기를 참조해 2026 월평균(N열)과 비교된다. 급여는 임원+직원+잡급, 차량유지비는 렌트료 포함, 잡손실은 처분손실·외환차손 포함으로 묶어 비교. 퇴직급여·도서인쇄비·내일채움공제는 우리 계정이 없어 비교 행 없음.', None),
    ('계정과목: 선택값 목록과 손익 구분 규칙. 계정을 추가하려면 이 시트에 행을 넣고 드롭다운 범위를 늘린다.', None),
    ('', None),
    ('[입력 규칙]', B),
    ('노란색 칸만 입력한다. 흰색 칸의 기존 확정 계정은 시스템에서 사용자가 확정한 값이며 필요하면 수정 가능. 수식 칸(검정·녹색 글자)은 건드리지 않는다.', None),
    ('필터·정렬은 자유롭게 써도 된다 — 수식(SUMIFS·COUNTIFS)은 숨겨진 행도 모두 집계한다. 단 필터가 걸린 상태에서 여러 행을 한 번에 붙여넣으면 숨겨진 행에 값이 들어갈 수 있으니, 붙여넣기 전에는 필터를 해제한다.', None),
    ('재출력(새 데이터 추가 후 다시 생성)해도 입력값은 유지된다 — 각 입력 시트 맨 오른쪽 숨김 열 "ID"(시스템 거래 ID)로 이전 파일의 입력을 자동으로 이어받는다. 열을 삭제하거나 행을 지우지 말 것. 파일은 그대로 다시 올리면 된다.', None),
    ('통장 출금 중 손익이 아닌 것: 계산서 대금 지급 = 2001 미지급금 / 카드대금 결제 = 2001 미지급금 / 대출 원금 상환 = 2002 단기차입금 또는 2004 장기차입금 / 내부 이체 = 0000 내부이체. 이자 납부만 5301 이자비용.', None),
    ('통장 입금 중 매출 입금 = 4001 상품매출 (손익에는 넣지 않음 — 매출은 홈택스 기준으로 대체). 대출 실행 입금 = 차입금 계정. 이자 수입 = 4002.', None),
    ('매입 계산서: 상품 매입 = 5002 상품매입(매출원가). 임차료·용역·비품 등은 해당 비용 계정 → 판매관리비로 간다. 미분류분은 손익현황에서 임시로 원가에 넣고 표시한다.', None),
    ('법인카드: 사용 건마다 비용 계정. 카드 결제대금이 통장에서 빠지는 출금은 2001 미지급금으로 처리해 이중 반영을 막는다.', None),
    ('이중 반영 방지: 계산서 대금으로 매칭된 통장 입출금(통장 시트 Q열 "계산서 매칭"=Y, 2026년 1,485건)은 계정과 무관하게 손익에서 자동 제외된다 — 비용·원가는 매입계산서 분류 시트가 담당. 매칭되지 않은 계산서 대금 출금은 사용자가 2001 미지급금으로 분류해야 중복이 없다(손익현황 참고 블록의 "매입 계산서 미매칭 공급가액"으로 확인).', None),
    ('통장 출금을 5002 상품매입으로 분류하면 매출원가에 들어간다 — 세금계산서를 받은 매입의 대금은 5002가 아니라 2001 미지급금이어야 한다(계산서가 이미 원가에 있어 이중 반영). 기존 확정값 97건이 5002로 되어 있으니 계산서 유무를 확인해 정리한다.', None),
    ('', None),
    ('[기준·가정]', B),
    ('손익은 공급가액(부가세 제외) 기준. 매출 = 매출 계산서 공급가(과세+면세) + 카드 매출 과세 공급가 + 카드 세액 0 금액(면세 또는 카드사 미분리) + 현금영수증 공급가. ERP 판매가는 부가세 포함이므로 참고 행만.', None),
    ('카드 매출 중 세액 0인 319건(2.4억)은 면세 상품일 수 있어 1.1로 나누지 않고 그대로 공급가로 둔다. 영업지원팀 확인 필요.', None),
    ('현금영수증 39건이 전량(2026-10-27 사용자 확인). 매출 계산서는 발행일 기준이라 9월 추석 매출의 10월 발행분은 9월에 없다.', None),
    ('월별손익현황에서 통장·카드·매입계산서 금액은 "계정과목(입력)" 값으로 SUMIFS 집계. 비용 = 출금 - 입금(환불 차감), 수익 = 입금 - 출금.', None),
    ('대출 이자는 통장 5301 분개가 실적이며, 대출관리현황의 월 이자는 계약 기준 계획으로 참고 행에만 둔다(이중 반영 없음).', None),
    ('데이터 출처: daolERP DB (transactions · card_expenses · tax_invoices · card_sales · cash_receipts · loans · accounts), 기간 2026-01-01 ~ 2026-09-30. ERP 참고값: 사용자 제공 2026 ERP 판매내역 파일.', None),
    ('', None),
    ('[예시] 통장 분개입력 시트 한 줄: 2026-03 | 2026-03-05 | 하나은행 보통 …6204 | 급여 3월 | 홍길동 | 출금 3,000,000 | 계정과목(입력) = 5101 급여 → 손익 구분(자동) = 판매관리비 → 월별손익현황 3월 급여 행에 합산.', None),
]
for i, (t, fnt) in enumerate(lines, 1):
    c = g.cell(i, 1, t); c.font = fnt or f(); c.alignment = Alignment(wrap_text=True, vertical='top')
g.cell(len(lines) + 2, 1, '범례:').font = B
c = g.cell(len(lines) + 3, 1, '노란색 = 입력·갱신할 칸'); c.fill = YELLOW; c.font = f()
c = g.cell(len(lines) + 4, 1, '녹색 글자 = 다른 시트에서 가져온 값(수식)'); c.font = GREEN
c = g.cell(len(lines) + 5, 1, '검정 글자 = 데이터 또는 같은 시트 수식'); c.font = f()

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and c.font.name != FONT:
                c.font = Font(name=FONT, size=c.font.size or 10, bold=c.font.bold, italic=c.font.italic, color=c.font.color)
import datetime
path = OUT + f'2026_월별손익현황_작성용_{datetime.date.today().isoformat()}.xlsx'
wb.save(path)
print('saved', path, 'tx rows', TX_LAST - 2, 'card rows', CD_LAST - 2, 'purchase rows', PI_LAST - 2)
print('HT rows', {k: v for k, v in HT.items() if '합계' in k}, 'PL rows', {k: v for k, v in PL.items() if '합계' in k or '이익' in k})
